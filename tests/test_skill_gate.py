import hashlib
import json
from types import SimpleNamespace

import openpyxl
import pytest

from benchmark_core.aggregation import aggregate
from benchmark_core.inspect_runner import run_aut
from benchmark_core.run_metrics import make_metrics
from benchmark_core.skill_gate import SkillGateError, enforce_or_block, evaluate_read_evidence
from benchmark_core.workspace import tree_hash


def make_config(backend='codex_cli'):
    return SimpleNamespace(data={'aut': {'skill_protocol': 'enforced_read_v1'},
                                 'docker_sandbox': {'execution_backend': backend}})


def fixture(name):
    text = f'---\nname: {name}\n---\nComplete contents\n'
    key = 'dns_skill_sha256' if name == 'kathara-dns' else 'creation_skill_sha256'
    return text, {key: hashlib.sha256(text.encode()).hexdigest()}


def read(name, content):
    return {'type': 'item.completed', 'item': {'id': name, 'type': 'command_execution',
            'command': f'cat .codex/skills/{name}/SKILL.md', 'exit_code': 0,
            'status': 'completed', 'aggregated_output': content}}


def evidence(tmp_path, mode, items, metadata=None, code=0, timed_out=False):
    path = tmp_path / 'logs/aut/events.jsonl'
    path.parent.mkdir(parents=True, exist_ok=True)
    events = items + [{'type': 'turn.completed'}]
    path.write_text('\n'.join(json.dumps(e) for e in events))
    return evaluate_read_evidence(tmp_path, {'skill_mode': mode, 'agent': 'codex', **(metadata or {})},
                                  events, returncode=code, timed_out=timed_out)


def test_evidence_full_partial_missing_wrong_and_timeout(tmp_path):
    content, meta = fixture('kathara-dns')
    assert evidence(tmp_path, 'dns_only', [read('kathara-dns', content)], meta)['blocked'] is False
    for items, m, timed_out in [([read('kathara-dns', content[:20])], meta, False),
                                ([], meta, False),
                                ([read('kathara-creation', content)], meta, False),
                                ([read('kathara-dns', content)], meta, True)]:
        assert evidence(tmp_path, 'dns_only', items, m, timed_out=timed_out)['blocked'] is True
    assert evidence(tmp_path, 'dns_only', [read('kathara-dns', content)], {}, code=0)['blocked'] is True


def test_evidence_both_requires_both(tmp_path):
    dns, dm = fixture('kathara-dns')
    creation, cm = fixture('kathara-creation')
    assert evidence(tmp_path, 'both_forced', [read('kathara-dns', dns)], {**dm, **cm})['blocked']
    verdict = evidence(tmp_path, 'both_forced', [read('kathara-dns', dns), read('kathara-creation', creation)], {**dm, **cm})
    assert not verdict['blocked'] and verdict['verified_skills'] == ['kathara-creation', 'kathara-dns']


@pytest.mark.parametrize('backend', ['codex_cli', 'docker_sandbox'])
def test_unsupported_backend_blocks_before_aut_and_preserves_lab(tmp_path, monkeypatch, backend):
    run = tmp_path / 'run'
    lab = run / 'lab'
    lab.mkdir(parents=True)
    (lab / 'lab.conf').write_text('original')
    before = tree_hash(lab)
    def must_not_run(*args, **kwargs):
        raise AssertionError('AUT was invoked')
    monkeypatch.setattr('benchmark_core.inspect_runner.run_codex', must_not_run)
    with pytest.raises(SkillGateError, match='unsupported_backend'):
        run_aut(make_config(backend), None, run, 'codex', 'original', skill_mode='dns_only')
    assert tree_hash(lab) == before
    record = json.loads((run / 'logs/aut/skill_gate.json').read_text())
    assert record['skill_gate_blocked'] is True and record['skill_gate_status'] == 'unsupported_backend'
    assert record['skill_gate_write_attempts_blocked'] is None


@pytest.mark.parametrize('mode', ['auto', 'no_skill'])
def test_no_gate_for_optional_modes(tmp_path, mode):
    enforce_or_block(make_config(), tmp_path, 'codex', mode)
    data = json.loads((tmp_path / 'logs/aut/skill_gate.json').read_text())
    assert data['skill_gate_enabled'] is False and data['skill_gate_blocked'] is False
    assert data['skill_gate_status'] == 'not_required'


def test_gate_reports_json_csv_excel(tmp_path):
    runs = tmp_path / 'runs'
    run = runs / 'lab/T1/dns_only/r001'
    run.mkdir(parents=True)
    with pytest.raises(SkillGateError):
        enforce_or_block(make_config(), run, 'codex', 'dns_only')
    manifest = {'agent': 'codex', 'skill_mode': 'dns_only', 'skill_protocol': 'enforced_read_v1', 'scenario_id': 'lab', 'prompt_type': 'T1'}
    run.joinpath('manifest.json').write_text(json.dumps(manifest))
    data = make_metrics(run, manifest, total_seconds=2, checker_seconds=None, checker_outcome=None)
    assert data['skill_gate_status'] == 'unsupported_backend' and data['skill_gate_input_tokens'] is None
    assert data['skill_gate_protocol_version'] == 'enforced_read_v1'
    assert data['skill_gate_isolation_verified'] is None and data['skill_gate_write_seconds'] is None
    assert data['skill_gate_state_history'][0]['status'] == 'unsupported_backend'
    (run / 'evaluation').mkdir()
    (run / 'evaluation/metrics.json').write_text(json.dumps(data))
    rows, _, summary = aggregate(runs, tmp_path / 'results')
    assert bool(rows.iloc[0]['skill_gate_blocked']) is True
    assert summary.iloc[0]['skill_gate_unsupported_runs'] == 1
    assert summary.iloc[0]['skill_gate_pass_denominator'] == 1
    assert summary.iloc[0]['skill_gate_pass_rate'] == 0
    book = openpyxl.load_workbook(tmp_path / 'results/benchmark.xlsx', read_only=True)
    assert 'skill_gate_status' in [c.value for c in next(book['Runs'].iter_rows())]
    assert 'skill_gate_blocked_runs' in [c.value for c in next(book['Summary'].iter_rows())]
    assert 'skill_gate_pass_denominator' in [c.value for c in next(book['Summary'].iter_rows())]


def test_pipeline_fails_closed_before_checker(tmp_path, monkeypatch):
    from benchmark_core import pipeline
    from benchmark_core.config import Config
    from benchmark_core.scenario_loader import Scenario
    scenario_root = tmp_path / 'scenarios/test_lab'
    (scenario_root / 'lab').mkdir(parents=True)
    (scenario_root / 'lab/lab.conf').write_text('original')
    (scenario_root / 'correction.yaml').write_text('fixture')
    (tmp_path / 'prompt/test_lab').mkdir(parents=True)
    (tmp_path / 'prompt/test_lab/T1.md').write_text('Configure DNS')
    (tmp_path / 'skills/dns').mkdir(parents=True)
    (tmp_path / 'skills/dns/SKILL.md').write_text('---\nname: kathara-dns\n---\nDNS')
    config = Config(tmp_path, {'aut': {'version': 'test', 'dns_skill': 'skills/dns/SKILL.md', 'skill_protocol': 'enforced_read_v1'},
                               'sandbox': {'image': 'fixture'}, 'benchmark': {'timeout_seconds': 1}})
    monkeypatch.setattr(pipeline, 'component_versions', lambda: {})
    monkeypatch.setattr(pipeline, 'read_correction', lambda _: (b'fixture', 'sha'))
    def no_checker(*args, **kwargs):
        raise AssertionError('checker was invoked')
    monkeypatch.setattr(pipeline, 'run_checker', no_checker)
    run = pipeline.run_one(config, Scenario('test_lab', scenario_root), 'codex', 'T1', 'dns_only')
    manifest = json.loads((run / 'manifest.json').read_text())
    metrics = json.loads((run / 'evaluation/metrics.json').read_text())
    assert manifest['pipeline_state'] == 'AUT_FAILED' and manifest['skill_protocol'] == 'enforced_read_v1'
    assert metrics['skill_gate_status'] == 'unsupported_backend' and metrics['skill_gate_blocked'] is True
    assert (run / 'lab/lab.conf').read_text() == 'original'
    assert not (run / 'evaluation/correction.yaml').exists()


def test_gate_rejects_unproven_exit_and_turn_failure(tmp_path):
    content, meta = fixture('kathara-dns')
    without_exit = read('kathara-dns', content)
    del without_exit['item']['exit_code']
    assert evidence(tmp_path, 'dns_only', [without_exit], meta)['status'] == 'verification_failed'
    failed = read('kathara-dns', content)
    failed['item']['exit_code'] = 1
    assert evidence(tmp_path, 'dns_only', [failed], meta)['blocked'] is True
    path = tmp_path / 'logs/aut/events.jsonl'
    events = [read('kathara-dns', content), {'type': 'turn.completed'}, {'type': 'turn.failed'}]
    path.write_text('\n'.join(json.dumps(x) for x in events))
    assert evaluate_read_evidence(tmp_path, {'agent':'codex', 'skill_mode':'dns_only', **meta}, events, returncode=0, timed_out=False)['status'] == 'read_phase_failed'


def test_gate_rejects_changed_lab_and_wrong_hash(tmp_path):
    content, meta = fixture('kathara-dns')
    change = {'type': 'item.completed', 'item': {'id': 'write', 'type': 'file_change', 'status': 'completed'}}
    assert evidence(tmp_path, 'dns_only', [read('kathara-dns', content), change], meta)['status'] == 'isolation_violation'
    assert evidence(tmp_path, 'dns_only', [read('kathara-dns', content)], {'dns_skill_sha256': '0' * 64})['blocked'] is True
    assert evidence(tmp_path, 'dns_only', [read('kathara-dns', content[:-1])], meta)['blocked'] is True
