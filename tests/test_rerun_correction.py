import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import csv

from benchmark_core.config import Config
from benchmark_core.pipeline import reevaluate_run


class RerunCorrectionTest(unittest.TestCase):
    def fixture(self, root, mode="dns_only", number=1):
        run = root / f"runs/example_dns_001/T1/{mode}/r{number:03d}"
        for folder in ("lab", "input/lab", "logs/aut", "results/lab", "evaluation"):
            (run / folder).mkdir(parents=True, exist_ok=True)
        (run / "lab/lab.conf").write_text("lab original\n")
        (run / "input/lab/lab.conf").write_text("baseline\n")
        (run / "logs/aut/events.jsonl").write_text('{"type":"turn.completed","usage":{"total_tokens":17}}\n')
        (run / "logs/aut/result.json").write_text('{"duration_seconds":3}\n')
        (run / "results/results.csv").write_text("OLD COUNTS\n")
        (run / "results/lab/lab_result_all.csv").write_text("old tests\n")
        manifest = {"run_id": f"example_dns_001__T1__{mode}__r{number:03d}", "run_number": number,
                    "scenario_id": "example_dns_001", "skill_mode": mode, "agent": "codex",
                    "model": "test-model", "available_skills": [], "forced_skills": [],
                    "aut_execution_success": True, "checker_execution_success": True,
                    "task_success": True, "pipeline_state": "COMPLETED", "state_history": [],
                    "correction_sha256": "old"}
        (run / "manifest.json").write_text(json.dumps(manifest))
        old_metrics = {"scenario": "example_dns_001", "prompt_type": "T1", "skill_mode": mode, "run_number": number,
                       "run_id": manifest["run_id"], "agent": "codex", "model": "test-model",
                       "available_skills": [], "forced_skills": [], "status": "COMPLETED",
                       "agent_success": True, "tokens": {"total": 17},
                       "timing": {"agent_seconds": 3, "total_seconds": 4},
                       "checker": {"passed": 9, "failed": 0, "total": 9, "pass_rate": 1.0},
                       "correction_sha256": "old"}
        (run / "evaluation/metrics.json").write_text(json.dumps(old_metrics))
        canonical = root / "scenarios/example_dns_001/correction.yaml"
        canonical.parent.mkdir(parents=True, exist_ok=True)
        canonical.write_bytes(b"new correction bytes\n")
        return run, canonical

    def config(self, root):
        return Config(root, {"checker": {"no_cache": True}, "benchmark": {"timeout_seconds": 10},
                             "results": {"directory": "results"}})

    def test_in_place_snapshot_checker_metrics_and_aut_artifacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run, canonical = self.fixture(root)
            lab_before = (run / "lab/lab.conf").read_bytes()
            baseline_before = (run / "input/lab/lab.conf").read_bytes()
            aut_before = (run / "logs/aut/events.jsonl").read_bytes()
            outcome = {"checks_passed": 1, "checks_failed": 1, "checks_total": 2,
                       "check_pass_rate": 0.5, "task_success": False}
            with patch("benchmark_core.pipeline.read_correction", return_value=(canonical.read_bytes(), hashlib.sha256(canonical.read_bytes()).hexdigest())), \
                 patch("benchmark_core.pipeline.validate_correction", return_value=hashlib.sha256(canonical.read_bytes()).hexdigest()), \
                 patch("benchmark_core.pipeline.run_checker", return_value=outcome) as checker:
                reevaluate_run(self.config(root), run)
            checker.assert_called_once()
            self.assertEqual((run / "evaluation/correction.yaml").read_bytes(), canonical.read_bytes())
            manifest = json.loads((run / "manifest.json").read_text())
            metrics = json.loads((run / "evaluation/metrics.json").read_text())
            self.assertEqual(manifest["correction_sha256"], hashlib.sha256(canonical.read_bytes()).hexdigest())
            self.assertEqual(manifest["pipeline_state"], "COMPLETED")
            self.assertFalse(manifest["task_success"])
            self.assertEqual(metrics["checker"], {"passed": 1, "failed": 1, "total": 2, "pass_rate": 0.5})
            self.assertEqual(metrics["tokens"]["total"], 17)
            self.assertEqual((run / "lab/lab.conf").read_bytes(), lab_before)
            self.assertEqual((run / "input/lab/lab.conf").read_bytes(), baseline_before)
            self.assertEqual((run / "logs/aut/events.jsonl").read_bytes(), aut_before)
            self.assertFalse((run / "results/results.csv").exists())

    def test_checker_technical_failure_cannot_reuse_old_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run, canonical = self.fixture(root)
            with patch("benchmark_core.pipeline.read_correction", return_value=(canonical.read_bytes(), "new-sha")), \
                 patch("benchmark_core.pipeline.validate_correction", return_value="new-sha"), \
                 patch("benchmark_core.pipeline.run_checker", side_effect=RuntimeError("infrastructure")):
                reevaluate_run(self.config(root), run)
            manifest = json.loads((run / "manifest.json").read_text())
            metrics = json.loads((run / "evaluation/metrics.json").read_text())
            self.assertEqual(manifest["pipeline_state"], "CHECKER_FAILED")
            self.assertFalse(manifest["checker_execution_success"])
            self.assertIsNone(manifest["task_success"])
            self.assertEqual(metrics["checker"], {"passed": None, "failed": None, "total": None, "pass_rate": None})

    def test_invalid_correction_preserves_previous_evaluation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run, _ = self.fixture(root)
            previous = (run / "evaluation/metrics.json").read_bytes()
            results = (run / "results/results.csv").read_bytes()
            with patch("benchmark_core.pipeline.read_correction", side_effect=ValueError("invalid")), \
                 patch("benchmark_core.pipeline.run_checker") as checker:
                with self.assertRaisesRegex(ValueError, "invalid"):
                    reevaluate_run(self.config(root), run)
            checker.assert_not_called()
            self.assertEqual((run / "evaluation/metrics.json").read_bytes(), previous)
            self.assertEqual((run / "results/results.csv").read_bytes(), results)

    def test_cli_all_modes_selects_existing_runs_only_and_aggregates_once(self):
        from benchmark_core.scenario_loader import Scenario
        from scripts import run_benchmark

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            scenario_dir = root / "scenarios/example_dns_001"
            scenario_dir.mkdir(parents=True)
            prompt_dir = root / "prompt/example_dns_001"
            prompt_dir.mkdir(parents=True)
            (prompt_dir / "T1.md").write_text("prompt")
            scenario = Scenario("example_dns_001", scenario_dir)
            config = Config(root, {"aut": {"agent": "codex"}, "benchmark": {"repetitions": 1},
                                   "results": {"directory": "results"}})
            expected = []
            owner = self
            for mode in ("no_skill", "creation_only", "dns_only", "both_forced", "auto"):
                owner.fixture(root, mode, 1)
            def fake_reevaluate(_config, path):
                expected.append(path)
                return path
            with patch("sys.argv", ["run_benchmark.py", "--scenario", "example_dns_001", "--prompt-type", "T1", "--rerun-correction"]), \
                 patch.object(run_benchmark, "load_config", return_value=config), \
                 patch.object(run_benchmark, "discover_scenarios", return_value={scenario.scenario_id: scenario}), \
                 patch.object(run_benchmark, "preflight") as preflight, \
                 patch.object(run_benchmark, "run_one") as run_aut, \
                 patch.object(run_benchmark, "reevaluate_run", side_effect=fake_reevaluate), \
                 patch.object(run_benchmark, "aggregate") as aggregate:
                self.assertEqual(run_benchmark.main(), 0)
            preflight.assert_not_called()
            run_aut.assert_not_called()
            self.assertEqual([p.parent.name for p in expected],
                             ["no_skill", "creation_only", "dns_only", "both_forced", "auto"])
            self.assertTrue(all(p.name == "r001" for p in expected))
            self.assertEqual(aggregate.call_count, 1)

    def test_existing_workbook_is_rebuilt_from_new_canonical_artifacts(self):
        from benchmark_core.aggregation import aggregate

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run, canonical = self.fixture(root)
            reports = run / "results"
            (reports / "results.csv").write_text(
                'Student Name,Tests Passed,Tests Failed,Tests Total Number,Problems\n'
                'lab,0,1,1,old\n')
            (reports / "lab/lab_result_all.csv").write_text(
                'Test Description,Passed,Reason\nold check,False,old failure\n')
            (reports / "lab/lab_result_summary.csv").write_text(
                'Total Tests,Passed Tests,Failed\n1,0,1\n')
            (reports / "lab/lab_result_failed.csv").write_text(
                'Test Description,Passed,Reason\nold check,False,old failure\n')
            out = root / "published"
            aggregate(root / "runs", out)
            from openpyxl import load_workbook
            before = load_workbook(out / "benchmark.xlsx")
            self.assertEqual(before["Runs"].freeze_panes, "A2")
            self.assertEqual(before["Runs"]["A1"].fill.fgColor.rgb[-6:], "17365D")

            def new_checker(_config, _run, _snapshot):
                (reports / "results.csv").write_text(
                    'Student Name,Tests Passed,Tests Failed,Tests Total Number,Problems\n'
                    'lab,2,0,2,\n')
                (reports / "lab/lab_result_all.csv").write_text(
                    'Test Description,Passed,Reason\nnew check A,True,ok\nnew check B,True,ok\n')
                (reports / "lab/lab_result_summary.csv").write_text(
                    'Total Tests,Passed Tests,Failed\n2,2,0\n')
                (reports / "lab/lab_result_failed.csv").write_text(
                    'Test Description,Passed,Reason\n')
                return {"checks_passed": 2, "checks_failed": 0, "checks_total": 2,
                        "check_pass_rate": 1.0, "task_success": True}

            with patch("benchmark_core.pipeline.read_correction", return_value=(canonical.read_bytes(), "new-sha")), \
                 patch("benchmark_core.pipeline.validate_correction", return_value="new-sha"), \
                 patch("benchmark_core.pipeline.run_checker", side_effect=new_checker):
                reevaluate_run(self.config(root), run)
            aggregate(root / "runs", out)
            workbook = load_workbook(out / "benchmark.xlsx", data_only=True)
            runs = workbook["Runs"]
            headers = {cell.value: cell.column for cell in runs[1]}
            run_ids = [runs.cell(row, headers["run_id"]).value for row in range(2, runs.max_row + 1)
                       if runs.cell(row, headers["run_id"]).value]
            self.assertEqual(run_ids.count("example_dns_001__T1__dns_only__r001"), 1)
            self.assertEqual(runs.cell(2, headers["tests_passed"]).value, 2)
            self.assertEqual(runs.cell(2, headers["tests_total"]).value, 2)
            self.assertEqual(runs.cell(2, headers["pass_rate"]).value, 1.0)
            checks = workbook["Checks"]
            description_column = [cell.value for cell in checks[1]].index("test_description") + 1
            check_descriptions = [checks.cell(row, description_column).value for row in range(2, checks.max_row + 1)
                                  if checks.cell(row, description_column).value]
            self.assertEqual(check_descriptions, ["new check A", "new check B"])
            self.assertEqual(workbook["Summary"]["C2"].value, 1)
            self.assertEqual(workbook["Runs"].freeze_panes, "A2")
            self.assertEqual(workbook["Runs"]["A1"].fill.fgColor.rgb[-6:], "17365D")


if __name__ == "__main__":
    unittest.main()
