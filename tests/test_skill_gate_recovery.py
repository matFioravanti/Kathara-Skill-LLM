"""Deterministic orchestration tests; these do not attest real isolation."""
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from benchmark_core.skill_gate import run_recovery_gate, SkillGateError
from benchmark_core.workspace import write_json


class ScriptedSession:
    schedule = []
    truncate = False
    fail_at = None
    mutate_at = None
    instances = []

    def __init__(self, *, workspace, state, logs, config):
        self.workspace,self.state,self.logs=workspace,state,logs
        self.thread_id='same-thread'; self.cleanup=[]; self.commands=[]
        self.turns=[]; self.overhead=1; self.version='deterministic-test-double'
        self.creates=[]; self.prompts=[]; self.closed=0
        self.instances.append(self)

    def create(self, *, readonly):
        self.creates.append(readonly)

    def close(self):
        self.closed+=1
        self.cleanup.append({'ok':True})

    def turn(self, *, prompt, directory, timeout):
        directory.mkdir(parents=True)
        self.prompts.append(prompt)
        saved=prompt[:-20] if self.truncate and directory.name=='fallback' else prompt
        (directory/'prompt_sent.md').write_bytes(saved)
        (directory/'stderr.log').write_text('')
        events=[{'type':'thread.started','thread_id':self.thread_id}]
        index=len(self.turns)
        if directory.name=='phase_b':
            (self.workspace/'lab.conf').write_text('configured by test AUT')
        elif directory.name!='fallback':
            for name in self.schedule[index] if index<len(self.schedule) else []:
                raw=(self.workspace/'.codex/skills'/name/'SKILL.md').read_bytes().decode('utf-8')
                events.append({'type':'item.completed','item':{'id':f'item_{len(events)}','type':'command_execution',
                    'command':f'cat .codex/skills/{name}/SKILL.md','status':'completed','exit_code':0,'aggregated_output':raw}})
        events.append({'type':'turn.completed','usage':{'input_tokens':11,'output_tokens':3,'cached_input_tokens':5}})
        (directory/'events.jsonl').write_text(''.join(json.dumps(e)+'\n' for e in events))
        result={'returncode':0,'timed_out':False,'duration_seconds':2,'usage':{'input_tokens':11,'output_tokens':3,'cached_input_tokens':5},'stdin_delivery_verified':True,'session_id':self.thread_id}
        write_json(directory/'invocation.json',{'stdin_bytes':len(prompt),
                   'stdin_sha256':hashlib.sha256(prompt).hexdigest(),
                   'session_expected':None if directory.name=='phase_a' else self.thread_id})
        self.turns.append((directory,result)); write_json(directory/'result.json',result)
        if self.mutate_at: self.mutate_at(index)
        if self.fail_at==index: raise RuntimeError('Infrastructure failure')
        return events,result


@pytest.fixture
def setup(tmp_path):
    workspace=tmp_path/'workspace'; workspace.mkdir()
    (workspace/'lab.conf').write_text('baseline')
    root=tmp_path/'project'; root.mkdir()
    run=root/'runs/case/T1/both_forced/r001'; run.mkdir(parents=True)
    aut={'skill_protocol':'enforced_read_v1','skill_gate_recovery_policy':'recovery_v1'}
    metadata={'agent':'codex','skill_mode':'both_forced','skill_protocol':'enforced_read_v1'}
    for name,key in [('kathara-dns','dns_skill'),('kathara-creation','creation_skill')]:
        raw=f'---\nname: {name}\ndescription: Isolated test fixture.\n---\nComplete bytes for {name}.\r\n'.encode()
        path=root/'skills'/name/'SKILL.md'; path.parent.mkdir(parents=True); path.write_bytes(raw)
        staged=workspace/'.codex/skills'/name/'SKILL.md'; staged.parent.mkdir(parents=True); staged.write_bytes(raw)
        aut[key]=str(path.relative_to(root)); metadata[key+'_sha256']=hashlib.sha256(raw).hexdigest()
    write_json(run/'manifest.json',metadata)
    config=SimpleNamespace(root=root,data={'aut':aut,'benchmark':{'timeout_seconds':5},'docker_sandbox':{'execution_backend':'docker_sandbox'}},
                           path=lambda p:root/p,model=lambda:None,reasoning_effort=lambda:None)
    ScriptedSession.schedule=[]; ScriptedSession.truncate=False; ScriptedSession.fail_at=None
    ScriptedSession.mutate_at=None; ScriptedSession.instances=[]
    def execute(mode='both_forced'):
        return run_recovery_gate(config=config,run=run,workspace=workspace,original_prompt='ORIGINAL FUNCTIONAL PROMPT',mode=mode,session_factory=ScriptedSession)
    return config,run,workspace,execute


def record(run): return json.loads((run/'logs/aut/skill_gate.json').read_text())


@pytest.mark.parametrize('mode,skills',[('dns_only',['kathara-dns']),('creation_only',['kathara-creation']),('both_forced',['kathara-creation','kathara-dns'])])
def test_direct_read_opens_gate(setup,mode,skills):
    cfg,run,ws,execute=setup
    ScriptedSession.schedule=[skills]
    execute(mode)
    d=record(run); s=ScriptedSession.instances[0]
    assert d['skill_gate_status']=='completed'
    assert d['skill_gate_load_methods']==dict.fromkeys(skills,'read_verified')
    assert s.creates==[True,False] and s.prompts[-1]==b'ORIGINAL FUNCTIONAL PROMPT'
    assert (ws/'lab.conf').read_text()=='configured by test AUT'
    assert d['skill_gate_read_attempts']==1 and d['skill_gate_retry_count']==0
    assert d['skill_gate_initial_input_tokens']==11 and d['write_phase_input_tokens']==11


@pytest.mark.parametrize('schedule,attempts,method',[
    ([[],['kathara-dns']],2,'retry_verified'),
    ([[],[],['kathara-dns']],3,'retry_verified'),
    ([[],[],[]],3,'preloaded_fallback'),
])
def test_bounded_retry_and_fallback(setup,schedule,attempts,method):
    cfg,run,ws,execute=setup; ScriptedSession.schedule=schedule
    execute('dns_only'); d=record(run); s=ScriptedSession.instances[0]
    assert d['skill_gate_read_attempts']==attempts
    assert d['skill_gate_retry_count']==attempts-1
    assert d['skill_gate_load_methods']=={'kathara-dns':method}
    assert d['skill_gate_session_count']==1 and d['skill_gate_session_id']=='same-thread'
    assert len(list((run.parents[0]).glob('r*')))==1
    assert len(s.turns)<=5 and s.creates==[True,False] and s.closed>=2
    if method=='preloaded_fallback':
        prompt=(run/'logs/aut/fallback/prompt_sent.md').read_bytes()
        raw=cfg.path(cfg.data['aut']['dns_skill']).read_bytes()
        assert raw in prompt and b'\r\n' in prompt
        assert d['skill_gate_fallback_sha256_verified'] is True
        assert d['skill_gate_context_delivery_verified'] is True
        assert d['skill_gate_verified_skills']==['kathara-dns']
        assert d['skill_gate_retry_input_tokens']==22
        assert d['skill_gate_input_tokens']==44


def test_retry_only_missing_skill(setup):
    cfg,run,ws,execute=setup
    ScriptedSession.schedule=[['kathara-creation'],['kathara-dns']]
    execute(); s=ScriptedSession.instances[0]; d=record(run)
    assert b'kathara-creation' not in s.prompts[1]
    assert d['skill_gate_load_methods']=={'kathara-creation':'read_verified','kathara-dns':'retry_verified'}
    from benchmark_core.run_metrics import make_metrics
    meta=json.loads((run/'manifest.json').read_text())
    m=make_metrics(run,meta,total_seconds=10,checker_seconds=None,checker_outcome=None)
    assert set(m['skill_compliance']['observed_read_skills'])=={'kathara-creation','kathara-dns'}


def test_mixed_read_and_preload(setup):
    cfg,run,ws,execute=setup; ScriptedSession.schedule=[['kathara-creation'],[],[]]
    execute(); d=record(run)
    assert d['skill_gate_load_methods']=={'kathara-creation':'read_verified','kathara-dns':'preloaded_fallback'}
    assert d['skill_gate_fallback_skills']==['kathara-dns']
    assert b'BEGIN SKILL kathara-creation' not in (run/'logs/aut/fallback/prompt_sent.md').read_bytes()


def test_zero_retries_goes_to_fallback(setup):
    cfg,run,ws,execute=setup; cfg.data['aut']['skill_gate_max_retries']=0
    execute('dns_only'); d=record(run)
    assert d['skill_gate_read_attempts']==1 and d['skill_gate_retry_count']==0
    assert d['skill_gate_retry_input_tokens'] is None


@pytest.mark.parametrize('kind',['wrong_initial_hash','changed_source','truncated_fallback','changed_lab','infrastructure'])
def test_technical_failure_never_unlocks(setup,kind):
    cfg,run,ws,execute=setup
    if kind=='wrong_initial_hash':
        d=json.loads((run/'manifest.json').read_text()); d['dns_skill_sha256']='0'*64; write_json(run/'manifest.json',d)
    elif kind=='changed_source':
        ScriptedSession.mutate_at=staticmethod(lambda index:cfg.path(cfg.data['aut']['dns_skill']).write_bytes(b'changed') if index==2 else None)
    elif kind=='truncated_fallback': ScriptedSession.truncate=True
    elif kind=='changed_lab': ScriptedSession.mutate_at=staticmethod(lambda index:(ws/'lab.conf').write_text('illegal'))
    else: ScriptedSession.fail_at=0
    with pytest.raises((SkillGateError,RuntimeError)): execute('dns_only')
    d=record(run)
    assert d['skill_gate_status']=='technical_error' and d['skill_gate_blocked']
    assert not (run/'logs/aut/phase_b').exists()
    if ScriptedSession.instances: assert ScriptedSession.instances[0].creates==[True]


@pytest.mark.parametrize('value',[-1,11,True,'2'])
def test_invalid_retry_limits(setup,value):
    cfg,run,ws,execute=setup; cfg.data['aut']['skill_gate_max_retries']=value
    with pytest.raises(ValueError): execute()


def test_missing_usage_is_null_and_totals_not_cached_twice(setup):
    from benchmark_core.run_metrics import make_metrics
    cfg,run,ws,execute=setup; ScriptedSession.schedule=[['kathara-dns']]
    execute('dns_only')
    m=make_metrics(run,{'agent':'codex','skill_mode':'dns_only'},total_seconds=20,checker_seconds=None,checker_outcome=None)
    assert m['tokens']['input']==22 and m['tokens']['cached_input']==10
    assert m['tokens']['total']==28
    assert m['skill_gate_retry_input_tokens'] is None


def test_aggregation_separates_policy_and_exports_methods(setup,tmp_path):
    from benchmark_core.aggregation import aggregate
    import openpyxl
    cfg,run,ws,execute=setup; ScriptedSession.schedule=[['kathara-creation'],[],[]]
    execute()
    m={**record(run),'agent':'codex','skill_mode':'both_forced','skill_protocol':'enforced_read_v1','scenario':'case','prompt_type':'T1'}
    write_json(run/'evaluation/metrics.json',m)
    legacy=run.parent/'r002'; legacy.mkdir(); write_json(legacy/'manifest.json',json.loads((run/'manifest.json').read_text()))
    write_json(legacy/'evaluation/metrics.json',{**m,'skill_gate_recovery_policy_version':None,'skill_gate_load_methods':None})
    rows,_,summary=aggregate(run.parents[3],tmp_path/'reports')
    assert len(summary)==2
    recovery=summary[summary['skill_gate_recovery_policy_version']=='recovery_v1'].iloc[0]
    assert recovery['skill_gate_direct_runs']==1 and recovery['skill_gate_fallback_runs']==1
    assert recovery['skill_gate_retry_runs']==0 and recovery['skill_gate_delivery_denominator']==1
    assert recovery['skill_gate_fallback_rate']==1
    book=openpyxl.load_workbook(tmp_path/'reports/benchmark.xlsx',read_only=True)
    assert 'skill_gate_load_methods' in [c.value for c in next(book['Runs'].iter_rows())]


@pytest.mark.parametrize('mode,skills',[('dns_only',['kathara-dns']),('creation_only',['kathara-creation']),('both_forced',['kathara-creation','kathara-dns'])])
def test_ordinary_pipeline_reaches_checker(setup,monkeypatch,mode,skills):
    from benchmark_core import pipeline
    from benchmark_core.config import Config
    from benchmark_core.scenario_loader import Scenario
    cfg,_,_,_=setup
    source=cfg.root/'scenarios/tiny'; (source/'lab').mkdir(parents=True)
    (source/'lab/lab.conf').write_text('baseline')
    (source/'correction.yaml').write_text('fixture')
    (cfg.root/'prompt/tiny').mkdir(parents=True)
    (cfg.root/'prompt/tiny/T1.md').write_text('ORIGINAL FUNCTIONAL PROMPT')
    cfg.data['aut']['version']='test'
    cfg.data['sandbox']={'image':'fixture'}
    config=Config(cfg.root,cfg.data)
    ScriptedSession.schedule=[skills]
    monkeypatch.setattr('benchmark_core.docker_sandbox_runner.ProtectedCodexSession',ScriptedSession)
    monkeypatch.setattr(pipeline,'component_versions',lambda:{})
    monkeypatch.setattr(pipeline,'read_correction',lambda p:(b'fixture','sha'))
    monkeypatch.setattr(pipeline,'validate_correction',lambda p:'sha')
    checked=[]
    def checker(config,run,correction):
        assert (run/'lab/lab.conf').read_text()=='configured by test AUT'
        checked.append(run)
        return {'checks_passed':1,'checks_failed':0,'checks_total':1,'check_pass_rate':1,'task_success':True}
    monkeypatch.setattr(pipeline,'run_checker',checker)
    run=pipeline.run_one(config,Scenario('tiny',source),'codex','T1',mode)
    manifest=json.loads((run/'manifest.json').read_text())
    assert manifest['pipeline_state']=='COMPLETED' and len(checked)==1
    assert manifest['skill_gate_load_methods']==dict.fromkeys(skills,'read_verified')
    assert manifest['skill_gate_recovery_policy_version']=='recovery_v1'
    assert (source/'lab/lab.conf').read_text()=='baseline'


@pytest.mark.parametrize('kind',['ok','missing_thread','wrong_thread','failed_turn','privileged','timeout'])
def test_real_launcher_checks_thread_and_privilege_events(tmp_path,monkeypatch,kind):
    import subprocess
    from benchmark_core.docker_sandbox_runner import ProtectedCodexSession, SandboxError
    config=SimpleNamespace(model=lambda:None,reasoning_effort=lambda:None)
    session=ProtectedCodexSession(workspace=tmp_path,state=tmp_path/'state',logs=tmp_path/'logs',config=config)
    session.name='unit-only'; session.thread_id='expected-thread'
    status='NoNewPrivs:\t1\n'+''.join(f'{key}:\t0000000000000000\n' for key in ('CapEff','CapPrm','CapBnd','CapAmb'))
    events=[{'type':'skill_gate.process_attestation','uid':session.uid,'status':status}]
    if kind!='missing_thread': events.append({'type':'thread.started','thread_id':'other-thread' if kind=='wrong_thread' else 'expected-thread'})
    if kind=='privileged': events[0]['uid']=0
    events.append({'type':'turn.failed' if kind=='failed_turn' else 'turn.completed','usage':{'input_tokens':5,'output_tokens':1}})
    raw=''.join(json.dumps(e)+'\n' for e in events).encode()
    payload=b'Exact input bytes\r\nUnicode: \xc3\xa8\n'
    def fake_run(command,**kwargs):
        assert kwargs['input']==payload
        assert 'resume' in command and 'expected-thread' in command
        assert '--ephemeral' not in command and '--bounding-set=-all' in command
        if kind=='timeout': raise subprocess.TimeoutExpired(command,5,output=raw,stderr=b'timed out')
        return subprocess.CompletedProcess(command,0,raw,b'')
    monkeypatch.setattr(subprocess,'run',fake_run)
    directory=tmp_path/'logs/retries/attempt_01'
    if kind=='ok':
        _,result=session.turn(prompt=payload,directory=directory,timeout=5)
        assert result['stdin_delivery_verified'] and session.thread_id=='expected-thread'
    else:
        with pytest.raises(SandboxError): session.turn(prompt=payload,directory=directory,timeout=5)
        assert json.loads((directory/'result.json').read_text()).get('error')
    assert (directory/'prompt_sent.md').read_bytes()==payload


def recovery_metrics(run, mode):
    from benchmark_core.run_metrics import make_metrics
    metadata = json.loads((run/'manifest.json').read_text())
    metadata['skill_mode'] = mode
    return make_metrics(run, metadata, total_seconds=None, checker_seconds=None, checker_outcome=None)


@pytest.mark.parametrize('mode,schedule,expected_methods,autonomous', [
    ('dns_only', [['kathara-dns']], {'kathara-dns':'read_verified'}, True),
    ('dns_only', [[], ['kathara-dns']], {'kathara-dns':'retry_verified'}, True),
    ('dns_only', [[], [], []], {'kathara-dns':'preloaded_fallback'}, False),
    ('both_forced', [['kathara-creation','kathara-dns']],
        {'kathara-creation':'read_verified','kathara-dns':'read_verified'}, True),
    ('both_forced', [['kathara-creation'], [], []],
        {'kathara-creation':'read_verified','kathara-dns':'preloaded_fallback'}, False),
    ('both_forced', [[], [], []],
        {'kathara-creation':'preloaded_fallback','kathara-dns':'preloaded_fallback'}, False),
])
def test_recovery_conformity_dimensions_per_skill(setup,mode,schedule,expected_methods,autonomous):
    _,run,_,execute=setup
    ScriptedSession.schedule=schedule
    execute(mode)
    metrics=recovery_metrics(run,mode)
    assert metrics['skill_autonomous_read_before_write'] is autonomous
    assert metrics['skill_delivery_before_write'] is True
    assert set(metrics['skill_delivery_details']) == set(expected_methods)
    for skill, method in expected_methods.items():
        assert metrics['skill_delivery_details'][skill] == {
            'method':method, 'autonomous_read_before_write':method != 'preloaded_fallback',
            'delivery_before_write':True}
    compliance=metrics['skill_compliance']
    assert compliance['required_reads_complete_before_modification'] is autonomous
    assert compliance['protocol_compliance_observed'] is True
    assert compliance['strict_status']==('verified' if autonomous else 'fallback_delivered')
    assert compliance['required_complete_read_count']==sum(m != 'preloaded_fallback' for m in expected_methods.values())
    for skill, method in expected_methods.items():
        assert compliance['skill_details'][skill]['read_before_modification'] is (method != 'preloaded_fallback')
        if method == 'preloaded_fallback':
            assert compliance['skill_details'][skill]['complete_read_verified'] is None


def test_missing_or_inconsistent_gate_evidence_stays_unknown(setup):
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-dns']]
    execute('dns_only')
    gate=record(run)
    # A certification string, even alongside an observed read, cannot establish
    # a transition that has no matching lifecycle evidence.
    gate['skill_gate_state_history']=[{'status':'initial_read'}]
    write_json(run/'logs/aut/skill_gate.json',gate)
    m=recovery_metrics(run,'dns_only')
    assert m['skill_autonomous_read_before_write'] is None
    assert m['skill_delivery_before_write'] is None
    assert m['skill_delivery_details']['kathara-dns']['delivery_before_write'] is None
    assert m['skill_compliance']['strict_status']=='unverified'
    assert m['skill_compliance']['protocol_compliance_observed'] is None


def test_preunlock_technical_failure_is_not_delivery(setup):
    _,run,_,execute=setup
    ScriptedSession.fail_at=0
    with pytest.raises(RuntimeError): execute('dns_only')
    m=recovery_metrics(run,'dns_only')
    assert m['skill_gate_status']=='technical_error'
    assert m['skill_autonomous_read_before_write'] is None
    assert m['skill_delivery_before_write'] is None
    assert m['skill_delivery_details']['kathara-dns']['delivery_before_write'] is None


@pytest.mark.parametrize('kind',['partial','wrong_hash','missing_exit','wrong_session','truncated_fallback','bad_timestamp','missing_retry','missing_b','wrong_protocol','contradictory_fallback'])
def test_corrupt_or_contradictory_evidence_never_proves_delivery(setup,kind):
    _,run,_,execute=setup
    ScriptedSession.schedule=([[],[],[]] if kind=='truncated_fallback' else [[],['kathara-dns']] if kind=='missing_retry' else [['kathara-dns']])
    execute('dns_only')
    metadata=json.loads((run/'manifest.json').read_text()); metadata['skill_mode']='dns_only'
    if kind=='wrong_hash':
        metadata['dns_skill_sha256']='0'*64
    elif kind=='wrong_session':
        path=run/'logs/aut/phase_b/result.json'; data=json.loads(path.read_text()); data['session_id']='other'; write_json(path,data)
    elif kind=='truncated_fallback':
        path=run/'logs/aut/fallback/prompt_sent.md'; path.write_bytes(path.read_bytes()[:-10])
    elif kind=='bad_timestamp':
        gate=record(run)
        passed=next(x for x in gate['skill_gate_state_history'] if x['status']=='passed')
        passed['timestamp']='9999-01-01T00:00:00+00:00'
        write_json(run/'logs/aut/skill_gate.json',gate)
    elif kind=='missing_retry':
        (run/'logs/aut/retries/attempt_01').rename(run/'logs/aut/missing_retry')
    elif kind=='missing_b':
        (run/'logs/aut/phase_b/result.json').rename(run/'logs/aut/phase_b/missing_result.json')
    elif kind in ('wrong_protocol','contradictory_fallback'):
        gate=record(run)
        if kind=='wrong_protocol': gate['skill_gate_protocol_version']='explicit_read_v2'
        else: gate['skill_gate_fallback_skills']=['kathara-dns']
        write_json(run/'logs/aut/skill_gate.json',gate)
    else:
        path=run/'logs/aut/phase_a/events.jsonl'; events=[json.loads(line) for line in path.read_text().splitlines()]
        read=next(e['item'] for e in events if e.get('item',{}).get('type')=='command_execution')
        if kind=='partial': read['aggregated_output']=read['aggregated_output'][:20]
        else: del read['exit_code']
        path.write_text(''.join(json.dumps(e)+'\n' for e in events))
    from benchmark_core.run_metrics import make_metrics
    m=make_metrics(run,metadata,total_seconds=None,checker_seconds=None,checker_outcome=None)
    assert m['skill_autonomous_read_before_write'] is None
    assert m['skill_delivery_before_write'] is None
    assert m['skill_compliance']['strict_status']=='unverified'
    assert m['skill_compliance']['skill_details']['kathara-dns']['complete_read_verified'] is None


def test_phase_b_read_cannot_be_counted_as_prior_read(setup):
    cfg,run,_,execute=setup
    execute('dns_only')
    path=run/'logs/aut/phase_b/events.jsonl'
    events=[json.loads(line) for line in path.read_text().splitlines()]
    events.insert(1,{'type':'item.completed','item':{'id':'extra','type':'command_execution',
        'command':'cat .codex/skills/kathara-dns/SKILL.md','status':'completed','exit_code':0,
        'aggregated_output':cfg.path(cfg.data['aut']['dns_skill']).read_bytes().decode('utf-8')}})
    path.write_text(''.join(json.dumps(e)+'\n' for e in events))
    m=recovery_metrics(run,'dns_only')
    assert m['skill_autonomous_read_before_write'] is False
    assert m['skill_delivery_before_write'] is True
    assert m['skill_compliance']['required_complete_read_count']==0
    assert m['skill_compliance']['observed_read_skills']==[]


def test_non_recovery_modes_keep_unknown_delivery(setup):
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-dns']]; execute('dns_only')
    metadata=json.loads((run/'manifest.json').read_text())
    from benchmark_core.run_metrics import make_metrics
    from benchmark_core.run_metrics import load_jsonl
    from benchmark_core.aut_observation import observe
    for mode,protocol in [('dns_only','explicit_read_v2'),('no_skill','enforced_read_v1'),('auto','enforced_read_v1')]:
        current={**metadata,'skill_mode':mode,'skill_protocol':protocol}
        m=make_metrics(run,current,total_seconds=None,checker_seconds=None,checker_outcome=None)
        assert m['skill_autonomous_read_before_write'] is None
        assert m['skill_delivery_before_write'] is None
        assert m['skill_delivery_details'] is None
        assert m['skill_compliance']==observe(run,current,load_jsonl(run/'logs/aut/events.jsonl'))[1]


def test_delivery_is_false_only_for_a_demonstrated_violation(setup):
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-dns']]; execute('dns_only')
    path=run/'logs/aut/phase_a/events.jsonl'
    events=[json.loads(line) for line in path.read_text().splitlines()]
    events=[e for e in events if e.get('item',{}).get('type')!='command_execution']
    path.write_text(''.join(json.dumps(e)+'\n' for e in events))
    gate=record(run)
    gate.update(skill_gate_context_delivery_verified=False,skill_gate_certification='FAILED',
                skill_gate_verified_skills=[],skill_gate_load_methods={},skill_gate_status='technical_error')
    write_json(run/'logs/aut/skill_gate.json',gate)
    m=recovery_metrics(run,'dns_only')
    assert m['skill_delivery_before_write'] is False
    assert m['skill_autonomous_read_before_write'] is None
    assert m['skill_compliance']['protocol_compliance_observed'] is False
    assert m['skill_compliance']['strict_status']=='delivery_failed'
    # Missing/truncated evidence is not proof that delivery was absent.
    (run/'logs/aut/phase_a/events.jsonl').write_text('{incomplete')
    m=recovery_metrics(run,'dns_only')
    assert m['skill_delivery_before_write'] is None


def test_aggregate_delivery_denominators_and_csv_excel(setup,tmp_path):
    import copy
    import csv
    import openpyxl
    from benchmark_core.aggregation import aggregate
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-creation','kathara-dns']]; execute()
    template=recovery_metrics(run,'both_forced')
    root=tmp_path/'aggregate/runs'; reports=tmp_path/'aggregate/results'
    situations=[('direct',True,True,'read_verified','completed'),
                ('retry',True,True,'retry_verified','completed'),
                ('fallback',False,True,'preloaded_fallback','completed'),
                ('unknown',None,None,None,'pending'),
                ('error',None,None,None,'technical_error'),
                ('checker_error',True,True,'read_verified','completed'),
                ('agent_error',True,True,'read_verified','technical_error'),
                ('correction_error',True,True,'read_verified','completed')]
    for index,(name,autonomous,delivery,method,status) in enumerate(situations,1):
        target=root/'case/T1/both_forced'/f'r{index:03d}'; target.mkdir(parents=True)
        data=copy.deepcopy(template)
        details={}
        for skill in ('kathara-creation','kathara-dns'):
            item_method = 'read_verified' if name=='fallback' and skill=='kathara-creation' else method
            skill_autonomous = True if item_method in ('read_verified','retry_verified') else False if item_method=='preloaded_fallback' else None
            details[skill]={'method':item_method,'autonomous_read_before_write':skill_autonomous,'delivery_before_write':delivery}
        data.update(run_id=name,scenario='case',prompt_type='T1',skill_mode='both_forced',
                    experiment_id='synthetic',skill_autonomous_read_before_write=autonomous,
                    skill_delivery_before_write=delivery,skill_delivery_details=details,
                    skill_gate_status=status,skill_gate_load_methods={k:v['method'] for k,v in details.items()},
                    status='CHECKER_FAILED' if name=='checker_error' else 'CORRECTION_INVALID' if name=='correction_error' else 'AUT_FAILED' if name in ('error','agent_error') else 'COMPLETED' if status=='completed' else 'PENDING')
        if name=='error':
            data['skill_gate_state_history']=[{'status':'preparing'},{'status':'initial_read'},{'status':'technical_error'}]
        elif name=='agent_error':
            data['skill_gate_state_history']=[{'status':'preparing'},{'status':'passed'},{'status':'write_phase_running'},{'status':'technical_error'}]
        write_json(target/'manifest.json',{'agent':'codex','skill_protocol':'enforced_read_v1','skill_mode':'both_forced'})
        write_json(target/'evaluation/metrics.json',data)
    rows,_,summary=aggregate(root,reports)
    assert len(rows)==8 and len(summary)==1
    s=summary.iloc[0]
    assert (s['skill_conformity_eligible_runs'],s['skill_autonomous_read_verified_runs'],s['skill_autonomous_read_denominator'])==(8,5,6)
    assert s['skill_autonomous_read_rate']==pytest.approx(5/6)
    assert (s['strict_verified_runs'],s['strict_compliance_denominator'])==(5,6)
    assert s['strict_observed_compliance_rate']==pytest.approx(5/6)
    assert (s['skill_delivery_verified_runs'],s['skill_delivery_denominator'],s['skill_delivery_rate'])==(6,6,1)
    assert (s['skill_fallback_verified_runs'],s['skill_fallback_denominator'])==(1,6)
    assert s['skill_fallback_rate']==pytest.approx(1/6)
    assert (s['skill_conformity_unknown_runs'],s['skill_technical_error_runs'])==(1,1)
    assert s['skill_conformity_unknown_rate']==pytest.approx(1/8)
    assert s['skill_technical_error_rate']==pytest.approx(1/8)
    assert (s['skill_gate_error_runs'],s['agent_post_gate_error_runs'],s['checker_error_runs'],s['correction_error_runs'])==(1,1,1,1)
    skills=json.loads(s['skill_delivery_by_skill'])
    assert skills['kathara-creation']['autonomous_read_verified_runs']==6
    assert skills['kathara-dns']['fallback_verified_runs']==1
    assert skills['kathara-dns']['unknown_runs']==1 and skills['kathara-dns']['technical_error_runs']==1
    assert (skills['kathara-dns']['agent_post_gate_error_runs'],skills['kathara-dns']['checker_error_runs'],skills['kathara-dns']['correction_error_runs'])==(1,1,1)
    with (reports/'runs.csv').open(newline='') as f:
        csv_rows={r['ID']:r for r in csv.DictReader(f)}
    assert csv_rows['fallback']['skill_autonomous_read_before_write']=='False'
    assert csv_rows['fallback']['skill_delivery_before_write']=='True'
    assert csv_rows['unknown']['skill_delivery_before_write']==''
    assert json.loads(csv_rows['fallback']['skill_delivery_details'])['kathara-dns']['method']=='preloaded_fallback'
    book=openpyxl.load_workbook(reports/'benchmark.xlsx',read_only=True)
    for name in ('Runs','Summary'):
        headers=[cell.value for cell in next(book[name].iter_rows())]
        assert ('skill_delivery_details' if name=='Runs' else 'skill_delivery_by_skill') in headers
    sheet=book['Runs']; headers=[cell.value for cell in next(sheet.iter_rows())]
    excel_rows={r[headers.index('ID')]:r for r in sheet.iter_rows(min_row=2,values_only=True)}
    assert excel_rows['fallback'][headers.index('skill_autonomous_read_before_write')] is False
    assert excel_rows['unknown'][headers.index('skill_delivery_before_write')] is None
    stored=json.loads((root/'case/T1/both_forced/r003/evaluation/metrics.json').read_text())
    assert stored['skill_delivery_details']['kathara-dns']['method']=='preloaded_fallback'


def test_partial_evidence_preserves_verified_individual_skill(setup):
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-creation','kathara-dns']]
    execute()
    path=run/'logs/aut/phase_a/events.jsonl'
    events=[json.loads(line) for line in path.read_text().splitlines()]
    dns=next(e['item'] for e in events if 'kathara-dns/SKILL.md' in e.get('item',{}).get('command',''))
    dns['aggregated_output']=dns['aggregated_output'][:10]
    path.write_text(''.join(json.dumps(e)+'\n' for e in events))
    m=recovery_metrics(run,'both_forced')
    assert m['skill_delivery_before_write'] is None
    assert m['skill_autonomous_read_before_write'] is None
    assert m['skill_delivery_details']['kathara-creation']['autonomous_read_before_write'] is True
    assert m['skill_delivery_details']['kathara-dns']['autonomous_read_before_write'] is None
    assert m['skill_compliance']['skill_details']['kathara-creation']['complete_read_verified'] is True
    assert m['skill_compliance']['skill_details']['kathara-dns']['complete_read_verified'] is None


def test_historical_strict_denominator_is_unchanged(tmp_path):
    from benchmark_core.aggregation import aggregate
    root=tmp_path/'runs'
    for i,value in enumerate((True,False,None),1):
        run=root/'legacy/T1/dns_only'/f'r{i:03d}'
        write_json(run/'manifest.json',{'agent':'codex','scenario_id':'legacy','prompt_type':'T1','skill_mode':'dns_only','skill_protocol':'explicit_read_v2'})
        write_json(run/'evaluation/metrics.json',{'agent':'codex','scenario':'legacy','prompt_type':'T1','skill_mode':'dns_only',
            'skill_protocol':'explicit_read_v2','skill_compliance':{'required_skills':['kathara-dns'],'protocol_compliance_observed':value}})
    _,_,summary=aggregate(root,tmp_path/'results')
    assert len(summary)==1
    row=summary.iloc[0]
    assert row['strict_verified_runs']==1 and row['strict_compliance_denominator']==3
    assert row['strict_observed_compliance_rate']==pytest.approx(1/3)


@pytest.mark.parametrize('history',[None,[{'status':'passed'},{'status':'write_phase_running'},{'status':'completed'},{'status':'technical_error'}]])
def test_missing_gate_history_does_not_invent_error_stage(setup,tmp_path,history):
    from benchmark_core.aggregation import aggregate
    _,run,_,execute=setup
    ScriptedSession.schedule=[['kathara-dns']]; execute('dns_only')
    data=recovery_metrics(run,'dns_only')
    data.update(skill_gate_status='technical_error',status='AUT_FAILED',skill_gate_state_history=history)
    write_json(run/'evaluation/metrics.json',data)
    _,_,summary=aggregate(run.parents[3],tmp_path/'out')
    row=summary.iloc[0]
    assert row['skill_gate_error_runs']==0 and row['agent_post_gate_error_runs']==0
    assert row['skill_unclassified_error_runs']==1


def test_checker_and_correction_errors_are_counted_without_gate(tmp_path):
    from benchmark_core.aggregation import aggregate
    root=tmp_path/'runs'
    for index,status in enumerate(('CHECKER_FAILED','CORRECTION_INVALID'),1):
        run=root/'case/T1/no_skill'/f'r{index:03d}'
        write_json(run/'manifest.json',{'agent':'codex','scenario_id':'case','prompt_type':'T1','skill_mode':'no_skill','skill_protocol':'explicit_read_v2'})
        write_json(run/'evaluation/metrics.json',{'agent':'codex','scenario':'case','prompt_type':'T1','skill_mode':'no_skill','skill_protocol':'explicit_read_v2','status':status})
    _,_,summary=aggregate(root,tmp_path/'results')
    row=summary.iloc[0]
    assert row['checker_error_runs']==1 and row['correction_error_runs']==1
    assert row['skill_gate_error_runs']==0 and row['skill_technical_error_runs']==0
