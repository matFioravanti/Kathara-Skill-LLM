import json
from pathlib import Path

import openpyxl
import pytest

from benchmark_core.aut_observation import observe
from benchmark_core.aggregation import aggregate
from benchmark_core.skill_modes import SKILL_MODES, execution_prompt
from benchmark_core.run_metrics import make_metrics


def event(kind, ident, phase='completed', **kwargs):
    return {'type': 'item.' + phase, 'item': {'type': kind, 'id': ident, **kwargs}}


def observation(tmp_path, events, mode='both_forced', agent='codex', trace=True):
    path = tmp_path / 'logs/aut/events.jsonl'
    if trace:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text('\n'.join(json.dumps(e) for e in events))
    return observe(tmp_path, {'agent': agent, 'skill_mode': mode}, events)


def test_web_lifecycle_and_shell_attempt(tmp_path):
    events = [event('web_search', 'a', 'started'), event('web_search', 'a', result={'url': 'https://kathara.org/docs'}),
              event('command_execution', 'b', command='curl https://example.com', exit_code=7, status='failed')]
    internet, _ = observation(tmp_path, events)
    assert (internet['web_search_calls'], internet['web_search_attempts'], internet['shell_network_candidates']) == (1, 1, 1)
    assert internet['internet_used'] is True
    assert internet['search_result_domains'] == ['kathara.org'] and internet['visited_domains'] == []
    internet, _ = observation(tmp_path, events[2:])
    assert internet['internet_used'] is None and internet['status'] == 'network_attempt_unconfirmed'


def test_missing_partial_logs(tmp_path):
    internet, skill = observation(tmp_path, [], trace=False)
    assert internet['internet_used'] is None and internet['web_search_calls'] is None
    assert skill['required_reads_observed'] is None
    internet, skill = observation(tmp_path, [event('web_search', 'a', 'started')])
    assert internet['web_search_attempts'] == 1 and internet['internet_used'] is None
    assert skill['status'] == 'unknown'


@pytest.mark.parametrize('skill,reader', [('kathara-dns', 'cat'), ('kathara-creation', "sed -n '1,240p'")])
def test_read_skill(tmp_path, skill, reader):
    e = event('command_execution', 'a', command=f'{reader} .codex/skills/{skill}/SKILL.md', exit_code=0, status='completed', aggregated_output=f'---\nname: {skill}\n---')
    _, compliance = observation(tmp_path, [e])
    assert compliance['skill_states'][skill] == 'observed'
    assert compliance['required_reads_observed'] is None


def test_both_modes_and_antigravity(tmp_path):
    es = [event('command_execution', s, command=f'cat .codex/skills/{s}/SKILL.md', exit_code=0, aggregated_output=f'name: {s}') for s in ('kathara-dns', 'kathara-creation')]
    _, skill = observation(tmp_path, es)
    assert skill['required_reads_observed'] is True and skill['status'] == 'verified'
    for mode in ('auto', 'no_skill'):
        _, skill = observation(tmp_path, [], mode=mode)
        assert skill['required_skills'] == [] and skill['status'] == 'not_required'
    _, skill = observation(tmp_path, [], mode='no_skill', agent='antigravity')
    assert skill['required_skills'] == ['kathara-dns'] and skill['status'] == 'unknown'
    for mode in ('creation_only', 'dns_only', 'both_forced'):
        for name in SKILL_MODES[mode].forced_skills:
            assert '$' + name in execution_prompt(mode, 'prompt') and name + '/SKILL.md' in execution_prompt(mode, 'prompt')
    assert execution_prompt('auto', 'prompt') == execution_prompt('no_skill', 'prompt') == 'prompt'


def test_results_and_legacy(tmp_path):
    run = tmp_path / 'runs/lab/T1/dns_only/r001'
    (run / 'evaluation').mkdir(parents=True)
    manifest = {'agent': 'codex', 'skill_mode': 'dns_only', 'scenario_id': 'lab', 'prompt_type': 'T1'}
    (run / 'manifest.json').write_text(json.dumps(manifest))
    metrics = make_metrics(run, manifest, total_seconds=None, checker_seconds=None, checker_outcome={'checks_passed': 1, 'checks_failed': 0, 'checks_total': 1, 'check_pass_rate': 1})
    assert metrics['checker']['passed'] == 1 and metrics['skill_compliance']['status'] == 'unknown'
    (run / 'evaluation/metrics.json').write_text(json.dumps(metrics))
    runs, checks, summary = aggregate(tmp_path / 'runs', tmp_path / 'results')
    assert runs.iloc[0]['internet_used'] is None
    assert summary.iloc[0]['internet_unknown_runs'] == 1
    assert summary.iloc[0]['skill_unobserved_or_unknown_runs'] == 1
    wb = openpyxl.load_workbook(tmp_path / 'results/benchmark.xlsx')
    assert 'internet_used' in [c.value for c in wb['Runs'][1]]
    assert 'skill_compliance_rate' in [c.value for c in wb['Summary'][1]]
    assert wb['Runs'].cell(2, list(runs.columns).index('internet_used') + 1).value is None
    metrics.pop('internet'); metrics.pop('skill_compliance')
    (run / 'evaluation/metrics.json').write_text(json.dumps(metrics))
    runs, _, _ = aggregate(tmp_path / 'runs', tmp_path / 'results')
    assert runs.iloc[0]['internet_used'] is None


def test_backfill_rolls_back_on_aggregation_error(tmp_path, monkeypatch):
    import sys
    from scripts import backfill_observations as backfill
    run = tmp_path / 'runs/lab/T1/no_skill/r001'
    (run / 'evaluation').mkdir(parents=True)
    (run / 'manifest.json').write_text(json.dumps({'agent': 'codex', 'skill_mode': 'no_skill'}))
    path = run / 'evaluation/metrics.json'
    path.write_text('{"checker": {"passed": 1}}')
    before = path.read_bytes()
    output = tmp_path / 'results'
    output.mkdir()
    report = output / 'runs.csv'
    report.write_text('original')
    def broken(runs, results):
        report.write_text('partial')
        raise RuntimeError('failure')
    monkeypatch.setattr(backfill, 'aggregate', broken)
    monkeypatch.setattr(sys, 'argv', ['backfill', '--runs', str(tmp_path / 'runs'), '--output', str(output)])
    with pytest.raises(RuntimeError):
        backfill.main()
    assert path.read_bytes() == before and report.read_text() == 'original'


@pytest.mark.parametrize('command', [
    'echo "BIND validators not installed on host"',
    'printf "%s\\n" "configuration reviewed on host"',
    'echo "Host aliases outside zone files"',
    "rg 'host|curl' config.txt",
    '# curl https://example.com\necho ok',
])
def test_shell_false_positives(tmp_path, command):
    internet, _ = observation(tmp_path, [event('command_execution', 'x', command=command, exit_code=0)])
    assert internet['shell_network_candidates'] == 0
    assert internet['internet_used'] is None


@pytest.mark.parametrize('command', [
    'curl -fsS https://example.com | head -n 5',
    "printf '%s' ready && wget https://example.com/data || echo host",
    '/bin/zsh -lc "pwd && curl -fsS https://example.com | sed -n 1p"',
    'echo ready\nssh example.com',
])
def test_shell_external_commands(tmp_path, command):
    internet, _ = observation(tmp_path, [event('command_execution', 'x', command=command, exit_code=7, status='failed')])
    assert internet['shell_network_candidates'] == 1
    assert internet['internet_used'] is None


def test_local_dns_and_uncertain(tmp_path):
    events = [event('command_execution', 'a', command='dig @10.33.0.10 services.test'),
              event('command_execution', 'b', command='curl "$DYNAMIC_URL"'),
              event('command_execution', 'c', command='curl http://localhost:8080')]
    internet, _ = observation(tmp_path, events)
    assert internet['shell_network_candidates'] == 0
    assert internet['shell_network_local'] == 2
    assert internet['shell_network_uncertain'] == 1
    assert internet['internet_used'] is None


def test_results_domains_and_visits(tmp_path):
    # Minimal shape from the recorded Codex CLI, including an unrelated thumbnail.
    results = [{'url': 'https://GitHub.COM/org/repo', 'thumbnail_url': 'https://images.openai.com/img'},
               {'url': 'https://github.com/other'}, {'url': 'http://10.1.2.3/doc'},
               {'url': 'https://LOCALHOST:8080'}, {'url': 'https://kathara.org/a'}]
    events = [event('web_search', 'a', 'started', action={'type': 'search'}),
              event('web_search', 'a', action={'type': 'search'}, results=results),
              event('web_search', 'b', action={'type': 'open_page', 'url': 'https://github.com/'},
                    results=[{'url': 'https://github.com/org/repo'}]),
              event('web_search', 'c', action={'type': 'open_page', 'url': 'https://failed.example/'},
                    results=[{'title': 'Internal Error'}]),
              event('web_search', 'd', 'started', action={'type': 'search'})]
    internet, _ = observation(tmp_path, events)
    assert internet['web_search_calls'] == 1 and internet['web_search_attempts'] == 2
    assert internet['search_result_domains'] == ['github.com', 'kathara.org']
    assert internet['visited_domains'] == ['github.com']
    assert internet['external_domains'] == ['github.com']


def test_skill_failed_partial_and_native(tmp_path):
    target = '.codex/skills/kathara-dns/SKILL.md'
    events = [event('command_execution', 'a', command=f'cat {target}', status='failed', exit_code=1, aggregated_output='permission denied'),
              event('command_execution', 'b', command='echo host', aggregated_output='name: kathara-creation'),
              event('skill_invocation', 'c', name='kathara-creation', status='completed')]
    _, skill = observation(tmp_path, events)
    assert skill['skill_states']['kathara-dns'] == 'unknown'
    assert skill['skill_states']['kathara-creation'] == 'observed'
    events.append(event('command_execution', 'd', command=f'sed -n 1,3p {target} && false', status='failed', exit_code=1, aggregated_output='---\nname: kathara-dns\n'))
    _, skill = observation(tmp_path, events)
    assert skill['required_reads_observed'] is True


def test_antigravity_step_updates(tmp_path):
    # Exact envelope used by inspect_adapter: synthetic tool names and inputs.
    def step(ident, name, state, inputs, outputs=None):
        return {'event': 'step_update', 'step_update': {'step_type': 'tool', 'step_id': ident,
                'tool_name': name, 'state': state, 'inputs': inputs, 'outputs': outputs}}
    events = [step('1', 'web_search', 'RUNNING', {'query': 'kathara'}),
              step('1', 'web_search', 'COMPLETED', {'query': 'kathara'}, {'results': [{'url': 'https://kathara.org/'}]}),
              step('2', 'run_command', 'COMPLETED', {'command': 'curl https://example.com'}, 'ok'),
              step('3', 'view_file', 'COMPLETED', {'path': '/tmp/antigravity_skills/SKILL.md'}, {'content': '---\nname: kathara-dns\n---'})]
    internet, skill = observation(tmp_path, events, mode='no_skill', agent='antigravity')
    assert (internet['web_search_calls'], internet['web_search_attempts'], internet['shell_network_candidates']) == (1, 1, 1)
    assert internet['search_result_domains'] == ['kathara.org']
    assert skill['required_reads_observed'] is True


def test_summary_denominators_and_excel(tmp_path):
    runs_root = tmp_path / 'runs'
    for number, events in enumerate(([event('web_search', 's', results=[{'url': 'https://kathara.org'}])], [], None), 1):
        run = runs_root / f'lab/T1/dns_only/r{number:03}'
        (run / 'evaluation').mkdir(parents=True)
        manifest = {'agent': 'codex', 'skill_mode': 'dns_only', 'scenario_id': 'lab', 'prompt_type': 'T1', 'run_number': number}
        (run / 'manifest.json').write_text(json.dumps(manifest))
        if events is not None:
            logs = run / 'logs/aut'
            logs.mkdir(parents=True)
            logs.joinpath('events.jsonl').write_text('\n'.join(json.dumps(x) for x in events))
        metrics = make_metrics(run, manifest, total_seconds=None, checker_seconds=None, checker_outcome=None)
        run.joinpath('evaluation/metrics.json').write_text(json.dumps(metrics))
    runs, _, summary = aggregate(runs_root, tmp_path / 'results')
    row = summary.iloc[0]
    assert row['internet_confirmed_runs'] == 1 and row['internet_unknown_runs'] == 2
    assert row['internet_observed_rate'] == pytest.approx(1 / 3)
    assert row['web_search_calls_denominator'] == 1 and row['mean_web_search_calls'] == 1
    assert row['skill_compliance_denominator'] == 3 and row['skill_unobserved_or_unknown_runs'] == 3
    wb = openpyxl.load_workbook(tmp_path / 'results/benchmark.xlsx', read_only=True)
    headers = [c.value for c in next(wb['Runs'].iter_rows())]
    assert 'search_result_domains' in headers and 'visited_domains' in headers
    assert 'internet_observed_rate' in [c.value for c in next(wb['Summary'].iter_rows())]


def test_backfill_idempotent_and_preserves_checker(tmp_path, monkeypatch):
    import sys
    from scripts import backfill_observations as backfill
    run = tmp_path / 'runs/lab/T1/no_skill/r001'
    (run / 'evaluation').mkdir(parents=True)
    (run / 'manifest.json').write_text(json.dumps({'agent': 'codex', 'skill_mode': 'no_skill', 'scenario_id': 'lab', 'prompt_type': 'T1'}))
    path = run / 'evaluation/metrics.json'
    path.write_text(json.dumps({'checker': {'passed': 2}, 'tokens': {'total': 100}, 'timing': {'agent_seconds': 3}}))
    monkeypatch.setattr(sys, 'argv', ['backfill', '--runs', str(tmp_path / 'runs'), '--output', str(tmp_path / 'results')])
    backfill.main()
    first = {p: p.read_bytes() for p in [path, *(tmp_path / 'results').iterdir()]}
    backfill.main()
    assert first == {p: p.read_bytes() for p in first}
    data = json.loads(path.read_text())
    assert data['checker'] == {'passed': 2} and data['tokens'] == {'total': 100} and data['timing'] == {'agent_seconds': 3}


def test_optional_skill_and_prompt_mentions(tmp_path):
    target = '.codex/skills/kathara-dns/SKILL.md'
    events = [event('agent_message', 'a', text=f'Please read {target}'),
              event('command_execution', 'b', command=f'echo {target}', aggregated_output='name: kathara-dns', exit_code=0)]
    _, skill = observation(tmp_path, events, mode='auto')
    assert skill['observed_read_skills'] == [] and skill['status'] == 'not_required'
    events.append(event('command_execution', 'c', command=f'cat {target}', aggregated_output='name: kathara-dns', exit_code=0))
    _, skill = observation(tmp_path, events, mode='auto')
    assert skill['observed_read_skills'] == ['kathara-dns'] and skill['required_reads_observed'] is None


def test_mcp_and_string_results(tmp_path):
    events = [event('mcp_tool_call', 'x', server='web', tool='run', arguments={'search_query': [{'q': 'Kathara'}]}, result='https://KATHARA.org/docs https://kathara.org/other'),
              event('command_execution', 'y', command='curl https://example.com', status='completed', exit_code=0)]
    internet, _ = observation(tmp_path, events)
    assert internet['web_search_calls'] == 1 and internet['search_result_domains'] == ['kathara.org']
    internet, _ = observation(tmp_path, events[1:])
    assert internet['internet_used'] is None
