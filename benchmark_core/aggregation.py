"""Aggregazione deterministica dei soli artefatti già presenti nelle run."""
from pathlib import Path
from datetime import datetime
import os
import json
import re
import statistics
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .checker_runner import category as checker_category
from .checker_runner import checker_test_rows
from .skill_gate import FIELDS as SKILL_GATE_FIELDS

RECOVERY_MEAN_FIELDS = (
    'skill_gate_initial_read_seconds', 'skill_gate_retry_seconds', 'skill_gate_fallback_seconds',
    'skill_gate_initial_input_tokens', 'skill_gate_initial_output_tokens',
    'skill_gate_retry_input_tokens', 'skill_gate_retry_output_tokens',
    'skill_gate_fallback_input_tokens', 'skill_gate_fallback_output_tokens',
    'skill_gate_read_attempts', 'skill_gate_retry_count')

RUN_COLUMNS = [
    # 1. IDENTIFICAZIONE
    "ID", "Lab", "T",
    "experiment_id", "replicate", "run_number", "skill_mode", "skill_protocol",
    "agent", "model", "reasoning_effort",
    "llm_temperature", "llm_top_p",
    "sampling_source", "sampling_status", "sampling_transport", "sampling_collection_mode",
    "sampling_temperature_status", "sampling_top_p_status", "sampling_distinct_temperatures", "sampling_distinct_top_p",

    # 2. METRICHE
    "status", "pipeline_completed", "checker_executed", "task_success",
    "checks_passed", "checks_total", "pass_rate",
    "agent_seconds", "checker_seconds", "total_seconds", "cost",
    "prompt_sha256", "prompt_sent_sha256", "correction_sha256",
    "evaluation_revision", "last_reevaluated_at",

    # 3. SKILL
    "available_skills", "forced_skills",

    "internet_used", "internet_status", "web_search_calls", "web_search_attempts",
    "shell_network_candidates", "shell_network_local", "shell_network_uncertain",
    "search_result_domains", "visited_domains", "network_policy", "internet_trace_available",
    "required_skills", "observed_read_skills", "required_reads_observed", "skill_compliance_status",
    "required_skill_count", "required_read_attempt_count", "required_content_identified_count",
    "required_complete_read_count", "required_unverified_read_count",
    "required_reads_complete_before_modification", "protocol_compliance_observed", "skill_strict_status",
    "skill_autonomous_read_before_write", "skill_delivery_before_write", "skill_delivery_details",
    *SKILL_GATE_FIELDS,

    # 4. TOKEN
    "input_tokens", "cached_input_tokens", "output_tokens",
    "reasoning_tokens", "total_tokens",
]
CHECK_COLUMNS = [
    "ID", "Lab", "T", "experiment_id", "skill_mode", "repetition", "replicate",
    "agent", "model",
    "category", "test_description", "passed", "reason",
]
SUMMARY_COLUMNS = [
    "experiment_id", "Lab", "T", "skill_mode", "skill_protocol", "agent", "model", "reasoning_effort",
    "skill_gate_recovery_policy_version", "skill_gate_delivery_denominator",
    "skill_gate_direct_runs", "skill_gate_direct_rate",
    "skill_gate_retry_runs", "skill_gate_retry_rate", "skill_gate_fallback_runs", "skill_gate_fallback_rate",
    "skill_conformity_eligible_runs", "skill_autonomous_read_verified_runs",
    "skill_autonomous_read_denominator", "skill_autonomous_read_rate",
    "skill_delivery_verified_runs", "skill_delivery_denominator", "skill_delivery_rate",
    "skill_fallback_verified_runs", "skill_fallback_denominator", "skill_fallback_rate",
    "skill_conformity_unknown_runs", "skill_conformity_unknown_rate",
    "skill_technical_error_runs", "skill_technical_error_rate", "skill_delivery_by_skill",
    *(f'mean_{field}' for field in RECOVERY_MEAN_FIELDS),
    "internet_confirmed_runs", "internet_confirmed_denominator", "internet_confirmed_rate",
    "internet_observed_rate",
    "mean_web_search_calls", "web_search_calls_denominator", "internet_unknown_runs",
    "skill_required_runs", "skill_verified_runs", "skill_unobserved_or_unknown_runs",
    "skill_compliance_denominator", "skill_compliance_rate",
    "required_skill_count", "required_read_attempt_count", "required_content_identified_count",
    "required_complete_read_count", "required_unverified_read_count",
    "complete_before_modification_runs", "strict_verified_runs", "strict_compliance_denominator", "strict_observed_compliance_rate",
    "skill_gate_enabled_runs", "skill_gate_blocked_runs", "skill_gate_unsupported_runs",
    "skill_gate_passed_runs", "skill_gate_statuses", "skill_gate_required_skills", "skill_gate_verified_skills",
    "skill_gate_backend_enforcements", "skill_gate_write_attempts_blocked",
    "mean_skill_gate_read_seconds", "mean_skill_gate_verification_seconds",
    "mean_skill_gate_input_tokens", "mean_skill_gate_output_tokens",
    "skill_gate_pass_denominator", "skill_gate_pass_rate", "skill_gate_isolation_violations",
    "mean_skill_gate_write_seconds", "mean_skill_gate_total_seconds",
    "mean_write_phase_input_tokens", "mean_write_phase_output_tokens",
    "total_runs", "pipeline_completed_runs", "checker_executed_runs", "evaluated_runs",
    "sampling_observed_runs", "sampling_temperatures", "sampling_top_ps",
    "task_success_runs", "task_failed_runs", "task_success_rate",
    "mean_checks_passed", "mean_checks_total", "mean_pass_rate", "median_pass_rate", "stdev_pass_rate",
    "mean_agent_seconds", "median_agent_seconds", "stdev_agent_seconds",
    "mean_checker_seconds", "median_checker_seconds",
    "mean_total_seconds", "median_total_seconds", "stdev_total_seconds",
    "mean_input_tokens", "stdev_input_tokens", "mean_cached_input_tokens", "stdev_cached_input_tokens",
    "mean_output_tokens", "stdev_output_tokens",
    "mean_reasoning_tokens", "stdev_reasoning_tokens",
    "mean_total_tokens", "median_total_tokens", "stdev_total_tokens",
]

SKILL_MODE_ORDER = {name: index for index, name in enumerate(
    ("no_skill", "creation_only", "dns_only", "both_forced", "auto")
)}

SKILL_MODE_COLORS = {
    "no_skill": "F2F2F2",      # Grigio chiaro
    "creation_only": "DDEBF7", # Azzurro chiaro
    "dns_only": "FCE4D6",      # Arancione chiaro
    "both_forced": "E2EFDA",   # Verde chiaro
    "auto": "E4DFEC",          # Lilla/viola chiaro
}
PROMPT_ORDER = {f"T{index}": index for index in range(1, 7)}

# Fogli extra rimossi come richiesto. Vengono mantenuti solo Runs, Checks e Summary.


def _number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _read_json(path: Path) -> dict:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Il documento deve essere un oggetto JSON: {path}")
    return value


def _natural_key(value):
    return tuple((0, int(part)) if part.isdigit() else (1, part.casefold())
                 for part in re.split(r"(\d+)", str(value or "")))


def _sort_key(row: dict):
    try:
        repetition = int(row.get("repetition"))
    except (TypeError, ValueError):
        repetition = -1
    prompt = str(row.get("prompt_type") or "")
    return (
        _natural_key(row.get("experiment_id")), _natural_key(row.get("Lab")),
        PROMPT_ORDER.get(prompt, 99), prompt,
        SKILL_MODE_ORDER.get(row.get("skill_mode"), 99), str(row.get("skill_mode") or ""),
        str(row.get("skill_protocol") or ""),
        str(row.get("agent") or ""), str(row.get("model") or ""),
        str(row.get("reasoning_effort") or ""), repetition, str(row.get("ID") or ""),
    )


def _check_category(description: str) -> str:
    text = description.casefold()
    if "named is " in text:
        return "dns_authority"
    if any(service in text for service in ("bgpd is ", "ospfd is ", "ospf6d is ",
                                           "ripd is ", "zebra is ", "watchfrr is ")):
        return "routing"
    if "delegat" in text:
        return "dns_delegation"
    known = checker_category(description)
    if known == "dns_authority":
        return "dns_authority"
    if known == "local_ns":
        return "dns_resolution"
    if known == "dns_record":
        return "dns_record"
    if known == "http":
        return "http"
    if known == "reachability" or "reachable from device" in text:
        return "reachability"
    if known == "custom":
        if "route" in text or "routing" in text:
            return "routing"
        if "ip address" in text or "inet" in text or " addr " in f" {text} ":
            return "addressing"
        return "custom"
    if "resolv" in text or "name server" in text or "nameserver" in text:
        return "dns_resolution"
    if "dns" in text and any(word in text for word in ("record", "a record", "aaaa", "mx", "cname")):
        return "dns_record"
    if "dns" in text and any(word in text for word in ("authority", "authoritative")):
        return "dns_authority"
    if "dns" in text and "delegat" in text:
        return "dns_delegation"
    if any(word in text for word in ("checking the routing table", "route show", "routing table", "route ")):
        return "routing"
    if "reachable from device" in text or "ping" in text:
        return "reachability"
    if "ip address" in text or "inet6" in text or "interface" in text and "address" in text:
        return "addressing"
    if "check existence" in text or "collision domain" in text or "lab structure" in text:
        return "topology"
    if text.startswith(("http check", "http check on")):
        return "http"
    return "other"


def _mean(values):
    usable = [number for value in values if (number := _number(value)) is not None]
    return statistics.fmean(usable) if usable else None


def _median(values):
    usable = [number for value in values if (number := _number(value)) is not None]
    return statistics.median(usable) if usable else None


def _stdev(values):
    usable = [number for value in values if (number := _number(value)) is not None]
    return statistics.stdev(usable) if len(usable) > 1 else None


def _run_record(metrics: dict, manifest: dict) -> dict:
    timing = metrics.get("timing") or {}
    tokens = metrics.get("tokens") or {}
    checker = metrics.get("checker") or {}
    tests_passed = _number(checker.get("passed"))
    tests_total = _number(checker.get("total"))
    pass_rate = _number(checker.get("pass_rate"))
    if pass_rate is None and tests_passed is not None and tests_total:
        pass_rate = tests_passed / tests_total
    execution_flag = manifest.get("checker_execution_success")
    status = metrics.get("status") or manifest.get("pipeline_state")
    if "checker_execution_success" in manifest:
        checker_executed = execution_flag is not None
    else:
        checker_executed = (_number(timing.get("checker_seconds")) is not None or status == "COMPLETED")
    evaluated = (execution_flag is True or
                 ("checker_execution_success" not in manifest and status == "COMPLETED")) and tests_total is not None
    task_success = manifest.get("task_success")
    if task_success not in (True, False):
        failed_checks = _number(checker.get("failed"))
        task_success = (failed_checks == 0 if failed_checks is not None else tests_passed == tests_total) \
            if evaluated and tests_passed is not None else None

    # experiment_id: può essere None per run storiche senza questo campo
    experiment_id = metrics.get("experiment_id") or manifest.get("experiment_id")
    # replicate: retrocompatibile — se assente usa repetition/run_number
    replicate = (metrics.get("replicate") if isinstance(metrics.get("replicate"), int)
                 else manifest.get("replicate") if isinstance(manifest.get("replicate"), int)
                 else None)
    repetition = metrics.get("run_number", manifest.get("run_number"))
    if replicate is None:
        replicate = repetition  # retrocompatibilità

    def _skill_list_str(value) -> str | None:
        """Serializes a skill list to a semicolon-separated string, or None if unavailable/empty."""
        if isinstance(value, list):
            joined = ";".join(value)
            return joined if joined else None
        return None


    avail_raw = metrics.get("available_skills")
    if not isinstance(avail_raw, list):
        avail_raw = manifest.get("available_skills")
    forced_raw = metrics.get("forced_skills")
    if not isinstance(forced_raw, list):
        forced_raw = manifest.get("forced_skills")

    internet = metrics.get("internet") or {}
    compliance = metrics.get("skill_compliance") or {}
    llm_sampling = metrics.get("llm_sampling") or manifest.get("llm_sampling") or {}
    temperature = _number(llm_sampling.get("temperature"))
    top_p = _number(llm_sampling.get("top_p"))
    sampling_source = llm_sampling.get("source")
    sampling_status = llm_sampling.get("status")
    sampling_transport = llm_sampling.get("transport")

    return {
        **{field: (_skill_list_str(metrics.get(field)) if field in ("skill_gate_required_skills", "skill_gate_verified_skills", "skill_gate_fallback_skills") else json.dumps(metrics[field], ensure_ascii=False) if isinstance(metrics.get(field), (dict, list)) else metrics.get(field)) for field in SKILL_GATE_FIELDS},
        "experiment_id": experiment_id,
        "Lab": metrics.get("scenario") or manifest.get("scenario_id"),
        "T": metrics.get("prompt_type") or manifest.get("prompt_type"),
        "skill_mode": metrics.get("skill_mode") or manifest.get("skill_mode"),
        "skill_protocol": metrics.get("skill_protocol") or manifest.get("skill_protocol"),
        "repetition": repetition,
        "replicate": replicate,
        "agent": metrics.get("agent") or manifest.get("agent"),
        "model": metrics.get("model") or manifest.get("model"),
        "reasoning_effort": metrics.get("reasoning_effort") or manifest.get("reasoning_effort"),
        "llm_temperature": temperature,
        "llm_top_p": top_p,
        "sampling_source": sampling_source,
        "sampling_status": sampling_status,
        "sampling_transport": sampling_transport,
        "sampling_collection_mode": llm_sampling.get("collection_mode"),
        "sampling_temperature_status": llm_sampling.get("temperature_status"),
        "sampling_top_p_status": llm_sampling.get("top_p_status"),
        "sampling_distinct_temperatures": json.dumps(llm_sampling["distinct_temperatures"]) if "distinct_temperatures" in llm_sampling else None,
        "sampling_distinct_top_p": json.dumps(llm_sampling["distinct_top_p"]) if "distinct_top_p" in llm_sampling else None,
        "status": status,
        "pipeline_completed": manifest.get("pipeline_state", status) == "COMPLETED",
        "checker_executed": checker_executed,
        "task_success": task_success if isinstance(task_success, bool) else None,
        "checks_passed": tests_passed if evaluated else None,
        "checks_total": tests_total if evaluated else None,
        "pass_rate": pass_rate if evaluated else None,
        "agent_seconds": _number(timing.get("agent_seconds")),
        "checker_seconds": _number(timing.get("checker_seconds")),
        "total_seconds": _number(timing.get("total_seconds")),
        "input_tokens": _number(tokens.get("input")),
        "cached_input_tokens": _number(tokens.get("cached_input")),
        "output_tokens": _number(tokens.get("output")),
        "reasoning_tokens": _number(tokens.get("reasoning")),
        "total_tokens": _number(tokens.get("total")),
        "cost": _first_number(metrics.get("cost"), metrics.get("cost_usd"),
                               manifest.get("cost"), manifest.get("cost_usd")),
        "ID": metrics.get("run_id") or manifest.get("run_id"),
        "run_number": repetition,
        "prompt_sha256": metrics.get("prompt_sha256") or manifest.get("prompt_sha256"),
        "prompt_sent_sha256": metrics.get("prompt_sent_sha256") or manifest.get("prompt_sent_sha256"),
        "correction_sha256": metrics.get("correction_sha256") or manifest.get("correction_sha256"),
        "evaluation_revision": metrics.get("evaluation_revision", manifest.get("evaluation_revision", 0)),
        "last_reevaluated_at": metrics.get("last_reevaluated_at") or manifest.get("last_reevaluated_at"),
        "internet_used": internet.get("internet_used"),
        "internet_status": internet.get("status"),
        "web_search_calls": internet.get("web_search_calls"),
        "web_search_attempts": internet.get("web_search_attempts"),
        "shell_network_candidates": internet.get("shell_network_candidates"),
        "shell_network_local": internet.get("shell_network_local"),
        "shell_network_uncertain": internet.get("shell_network_uncertain"),
        "search_result_domains": _skill_list_str(internet.get("search_result_domains")),
        "visited_domains": _skill_list_str(internet.get("visited_domains")),
        "network_policy": internet.get("network_policy"),
        "internet_trace_available": internet.get("trace_available"),
        "required_skills": _skill_list_str(compliance.get("required_skills")),
        "observed_read_skills": _skill_list_str(compliance.get("observed_read_skills")),
        "required_reads_observed": compliance.get("required_reads_observed"),
        "skill_compliance_status": compliance.get("status"),
        "required_skill_count": compliance.get("required_skill_count"),
        "required_read_attempt_count": compliance.get("required_read_attempt_count"),
        "required_content_identified_count": compliance.get("required_content_identified_count"),
        "required_complete_read_count": compliance.get("required_complete_read_count"),
        "required_unverified_read_count": compliance.get("required_unverified_read_count"),
        "required_reads_complete_before_modification": compliance.get("required_reads_complete_before_modification"),
        "protocol_compliance_observed": compliance.get("protocol_compliance_observed"),
        "skill_strict_status": compliance.get("strict_status"),
        "skill_autonomous_read_before_write": metrics.get("skill_autonomous_read_before_write"),
        "skill_delivery_before_write": metrics.get("skill_delivery_before_write"),
        "skill_delivery_details": json.dumps(metrics['skill_delivery_details'], ensure_ascii=False, sort_keys=True)
                                  if isinstance(metrics.get('skill_delivery_details'), dict) else None,
        "available_skills": _skill_list_str(avail_raw),
        "forced_skills": _skill_list_str(forced_raw),

        "_evaluated": evaluated,
        "_manifest": manifest,
    }



def _first_number(*values):
    return next((number for value in values if (number := _number(value)) is not None), None)


def _check_records(run: Path, record: dict) -> list[dict]:
    try:
        raw_rows = checker_test_rows(run / "results")
    except (OSError, ValueError, UnicodeError):
        return []
    # The full report is the sole source. The checker also writes a failed-only report;
    # reading that file as well would duplicate failed checks.
    rows = []
    for check in raw_rows:
        description = check.get("test_description") or ""
        rows.append({
            "experiment_id": record.get("experiment_id"),
            "Lab": record.get("Lab"), "T": record.get("T"),
            "skill_mode": record.get("skill_mode"), "repetition": record.get("repetition"),
            "replicate": record.get("replicate"),
            "agent": record.get("agent"), "model": record.get("model"),
            "ID": record.get("ID"), "category": _check_category(description),
            "test_description": description, "passed": check.get("passed"), "reason": check.get("reason"),
        })
    return rows


def _summary_records(records: list[dict]) -> list[dict]:
    groups = {}
    group_fields = ("experiment_id", "Lab", "T", "skill_mode", "skill_protocol", "agent", "model", "reasoning_effort", "skill_gate_recovery_policy_version")
    for record in records:
        key = tuple(record.get(field) for field in group_fields)
        groups.setdefault(key, []).append(record)
    output = []
    for key, group in groups.items():
        evaluated = [row for row in group if row["_evaluated"]]
        task_successes = sum(row.get("task_success") is True for row in evaluated)
        task_failures = sum(row.get("task_success") is False for row in evaluated)
        pass_rates = [row.get("pass_rate") for row in evaluated]
        manifest_states = [row.get("pipeline_completed") for row in group]
        checker_seconds = [row.get("checker_seconds") for row in group]
        observed_sampling = [row for row in group if row.get("sampling_status") in ("observed", "partial")]
        observed_temps = sorted({
            row["llm_temperature"] for row in group if row.get("llm_temperature") is not None
            and (row.get("sampling_temperature_status") == "observed" or
                 (row.get("sampling_temperature_status") is None and row.get("sampling_status") == "observed"))
        })
        observed_top_ps = sorted({
            row["llm_top_p"] for row in group if row.get("llm_top_p") is not None
            and (row.get("sampling_top_p_status") == "observed" or
                 (row.get("sampling_top_p_status") is None and row.get("sampling_status") == "observed"))
        })
        sampling_temps_str = ";".join(str(t) for t in observed_temps) if observed_temps else None
        sampling_top_ps_str = ";".join(str(p) for p in observed_top_ps) if observed_top_ps else None
        stats = {}
        for column, values in (
            ("pass_rate", pass_rates),
            ("agent_seconds", [row.get("agent_seconds") for row in group]),
            ("checker_seconds", checker_seconds),
            ("total_seconds", [row.get("total_seconds") for row in group]),
            ("input_tokens", [row.get("input_tokens") for row in group]),
            ("cached_input_tokens", [row.get("cached_input_tokens") for row in group]),
            ("output_tokens", [row.get("output_tokens") for row in group]),
            ("reasoning_tokens", [row.get("reasoning_tokens") for row in group]),
            ("total_tokens", [row.get("total_tokens") for row in group]),
        ):
            stats[f"mean_{column}"] = _mean(values)
            stats[f"stdev_{column}"] = _stdev(values)
            if column in ("pass_rate", "agent_seconds", "checker_seconds", "total_seconds", "total_tokens"):
                stats[f"median_{column}"] = _median(values)
        
        skill_required = [row for row in group if row.get("required_skills")]
        skill_verified = sum(row.get("required_reads_observed") is True for row in skill_required)
        internet_confirmed = sum(row.get("internet_used") is True for row in group)
        strict_verified = sum(row.get("protocol_compliance_observed") is True for row in skill_required)
        gate_enabled = sum(row.get("skill_gate_enabled") is True for row in group)
        gate_passed = sum(row.get("skill_gate_status") in ("passed", "write_phase_running", "completed") for row in group)
        delivery = [row for row in group if row.get('skill_gate_enabled') is True
                    and row.get('skill_gate_recovery_policy_version') == 'recovery_v1']
        def delivery_details(row):
            try:
                data = json.loads(row['skill_delivery_details']) if isinstance(row.get('skill_delivery_details'), str) else row.get('skill_delivery_details')
                return data if isinstance(data, dict) else {}
            except ValueError:
                return {}
        def known_boolean(row, field):
            return type(row.get(field)) is bool
        autonomous_known = [row for row in delivery if known_boolean(row, 'skill_autonomous_read_before_write')]
        delivery_known = [row for row in delivery if known_boolean(row, 'skill_delivery_before_write')]
        autonomous_count = sum(row['skill_autonomous_read_before_write'] is True for row in autonomous_known)
        delivery_count = sum(row['skill_delivery_before_write'] is True for row in delivery_known)
        fallback_count = sum(any(d.get('method') == 'preloaded_fallback' and d.get('delivery_before_write') is True
                                 for d in delivery_details(row).values() if isinstance(d, dict)) for row in delivery_known)
        def technical_error(row):
            return (row.get('skill_gate_status') == 'technical_error'
                    or row.get('status') in ('AUT_FAILED', 'CHECKER_FAILED', 'CORRECTION_MISSING', 'CORRECTION_INVALID'))
        technical = [row for row in delivery if technical_error(row)]
        unknown = [row for row in delivery if not technical_error(row)
                   and (not known_boolean(row, 'skill_autonomous_read_before_write')
                        or not known_boolean(row, 'skill_delivery_before_write'))]
        by_skill = {}
        for row in delivery:
            entries = delivery_details(row)
            for skill in (row.get('skill_gate_required_skills') or '').split(';'):
                if not skill:
                    continue
                stats_for_skill = by_skill.setdefault(skill, {'required_runs': 0, 'autonomous_read_verified_runs': 0,
                    'autonomous_read_denominator': 0, 'delivery_verified_runs': 0, 'delivery_denominator': 0,
                    'fallback_verified_runs': 0, 'unknown_runs': 0, 'technical_error_runs': 0})
                detail = entries.get(skill) if isinstance(entries.get(skill), dict) else {}
                a, d = detail.get('autonomous_read_before_write'), detail.get('delivery_before_write')
                stats_for_skill['required_runs'] += 1
                stats_for_skill['autonomous_read_denominator'] += type(a) is bool
                stats_for_skill['autonomous_read_verified_runs'] += a is True
                stats_for_skill['delivery_denominator'] += type(d) is bool
                stats_for_skill['delivery_verified_runs'] += d is True
                stats_for_skill['fallback_verified_runs'] += detail.get('method') == 'preloaded_fallback' and d is True
                if technical_error(row):
                    stats_for_skill['technical_error_runs'] += 1
                elif type(a) is not bool or type(d) is not bool:
                    stats_for_skill['unknown_runs'] += 1
        for value in by_skill.values():
            for label, count, denominator in [('autonomous_read_rate', 'autonomous_read_verified_runs', 'autonomous_read_denominator'),
                                               ('delivery_rate', 'delivery_verified_runs', 'delivery_denominator'),
                                               ('fallback_rate', 'fallback_verified_runs', 'delivery_denominator')]:
                value[label] = value[count] / value[denominator] if value[denominator] else None
        def has_method(row, method):
            value = row.get('skill_gate_load_methods')
            if isinstance(value, str):
                value = json.loads(value)
            return isinstance(value, dict) and method in value.values()
        methods = {}
        for label, method in [('direct', 'read_verified'), ('retry', 'retry_verified'), ('fallback', 'preloaded_fallback')]:
            count = sum(has_method(row, method) for row in delivery)
            methods[f'skill_gate_{label}_runs'] = count
            methods[f'skill_gate_{label}_rate'] = count / len(delivery) if delivery else None
        def known_total(field):
            values = [_number(row.get(field)) for row in group]
            values = [v for v in values if v is not None]
            return sum(values) if values else None
        output.append({
            **dict(zip(group_fields, key)),
            **methods,
            'skill_conformity_eligible_runs': len(delivery),
            'skill_autonomous_read_verified_runs': autonomous_count,
            'skill_autonomous_read_denominator': len(autonomous_known),
            'skill_autonomous_read_rate': autonomous_count / len(autonomous_known) if autonomous_known else None,
            'skill_delivery_verified_runs': delivery_count,
            'skill_delivery_denominator': len(delivery_known),
            'skill_delivery_rate': delivery_count / len(delivery_known) if delivery_known else None,
            'skill_fallback_verified_runs': fallback_count,
            'skill_fallback_denominator': len(delivery_known),
            'skill_fallback_rate': fallback_count / len(delivery_known) if delivery_known else None,
            'skill_conformity_unknown_runs': len(unknown),
            'skill_conformity_unknown_rate': len(unknown) / len(delivery) if delivery else None,
            'skill_technical_error_runs': len(technical),
            'skill_technical_error_rate': len(technical) / len(delivery) if delivery else None,
            'skill_delivery_by_skill': json.dumps(by_skill, ensure_ascii=False, sort_keys=True) if by_skill else None,
            **{f'mean_{field}': _mean(row.get(field) for row in group) for field in RECOVERY_MEAN_FIELDS},
            'skill_gate_delivery_denominator': len(delivery),
            "internet_confirmed_runs": internet_confirmed,
            "internet_confirmed_denominator": len(group),
            "internet_confirmed_rate": internet_confirmed / len(group),
            "internet_observed_rate": internet_confirmed / len(group),
            "internet_unknown_runs": sum(row.get("internet_used") is None for row in group),
            "mean_web_search_calls": _mean(row.get("web_search_calls") for row in group),
            "web_search_calls_denominator": sum(_number(row.get("web_search_calls")) is not None for row in group),
            "skill_required_runs": len(skill_required),
            "skill_verified_runs": skill_verified,
            "skill_unobserved_or_unknown_runs": len(skill_required) - skill_verified,
            "skill_compliance_denominator": len(skill_required),
            "skill_compliance_rate": skill_verified / len(skill_required) if skill_required else None,
            **{field: known_total(field) for field in ("required_skill_count", "required_read_attempt_count", "required_content_identified_count", "required_complete_read_count", "required_unverified_read_count")},
            "complete_before_modification_runs": sum(row.get("required_reads_complete_before_modification") is True for row in skill_required),
            "skill_gate_enabled_runs": sum(row.get("skill_gate_enabled") is True for row in group),
            "skill_gate_blocked_runs": sum(row.get("skill_gate_blocked") is True for row in group),
            "skill_gate_unsupported_runs": sum(row.get("skill_gate_status") == "unsupported_backend" for row in group),
            "skill_gate_passed_runs": gate_passed,
            "skill_gate_pass_denominator": gate_enabled,
            "skill_gate_pass_rate": gate_passed / gate_enabled if gate_enabled else None,
            "skill_gate_isolation_violations": known_total("skill_gate_isolation_violations"),
            "skill_gate_statuses": ";".join(sorted({row["skill_gate_status"] for row in group if row.get("skill_gate_status")})) or None,
            "skill_gate_required_skills": ";".join(sorted({skill for row in group for skill in (row.get("skill_gate_required_skills") or "").split(";") if skill})) or None,
            "skill_gate_verified_skills": ";".join(sorted({skill for row in group for skill in (row.get("skill_gate_verified_skills") or "").split(";") if skill})) or None,
            "skill_gate_backend_enforcements": ";".join(sorted({row["skill_gate_backend_enforcement"] for row in group if row.get("skill_gate_backend_enforcement")})) or None,
            "skill_gate_write_attempts_blocked": known_total("skill_gate_write_attempts_blocked"),
            "mean_skill_gate_read_seconds": _mean(row.get("skill_gate_read_seconds") for row in group),
            "mean_skill_gate_verification_seconds": _mean(row.get("skill_gate_verification_seconds") for row in group),
            "mean_skill_gate_write_seconds": _mean(row.get("skill_gate_write_seconds") for row in group),
            "mean_skill_gate_total_seconds": _mean(row.get("skill_gate_total_seconds") for row in group),
            "mean_write_phase_input_tokens": _mean(row.get("write_phase_input_tokens") for row in group),
            "mean_write_phase_output_tokens": _mean(row.get("write_phase_output_tokens") for row in group),
            "mean_skill_gate_input_tokens": _mean(row.get("skill_gate_input_tokens") for row in group),
            "mean_skill_gate_output_tokens": _mean(row.get("skill_gate_output_tokens") for row in group),
            "strict_verified_runs": strict_verified,
            "strict_compliance_denominator": len(skill_required),
            "strict_observed_compliance_rate": strict_verified / len(skill_required) if skill_required else None,
            "total_runs": len(group),
            "pipeline_completed_runs": sum(state is True for state in manifest_states),
            "checker_executed_runs": sum(row.get("checker_executed") is True for row in group),
            "evaluated_runs": len(evaluated),
            "sampling_observed_runs": len(observed_sampling),
            "sampling_temperatures": sampling_temps_str,
            "sampling_top_ps": sampling_top_ps_str,
            "task_success_runs": task_successes,
            "task_failed_runs": task_failures,
            "task_success_rate": task_successes / len(evaluated) if evaluated else None,
            "mean_checks_passed": _mean(row.get("checks_passed") for row in evaluated),
            "mean_checks_total": _mean(row.get("checks_total") for row in evaluated),
            **stats,
        })

    return sorted(output, key=_sort_key)


# Fogli analitici rimossi (Dashboard, Matrix, Lab Summary, ecc.)



# ---------------------------------------------------------------------------
# Scrittura CSV e Excel
# ---------------------------------------------------------------------------

def _atomic_csv(path: Path, frame: pd.DataFrame) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    frame.to_csv(temporary, index=False, na_rep="", lineterminator="\n")
    temporary.replace(path)


def _style_sheet(sheet, frame: pd.DataFrame, is_runs_sheet: bool = False) -> None:
    # Colori per i tre blocchi nel foglio Runs
    header_font = Font(color="FFFFFF", bold=True)
    id_fill = PatternFill("solid", fgColor="2F5597")      # Blu scuro
    metrics_fill = PatternFill("solid", fgColor="385723") # Verde scuro
    tokens_fill = PatternFill("solid", fgColor="C55A11")  # Arancione scuro
    default_fill = PatternFill("solid", fgColor="17365D") # Colore di default per altri fogli
    
    body_fill = PatternFill("solid", fgColor="F4F7FA")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.auto_filter.ref = sheet.dimensions
    
    for column_index, column_name in enumerate(frame.columns, start=1):
        cell = sheet.cell(1, column_index)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        if is_runs_sheet:
            col_name_str = str(column_name)
            # Determina il gruppo in base all'indice o nome
            if "token" in col_name_str.lower():
                cell.fill = tokens_fill
            elif col_name_str in (
                "ID", "Lab", "T", "experiment_id", "replicate", "run_number",
                "skill_mode", "skill_protocol", "agent", "model", "reasoning_effort",
                "llm_temperature", "llm_top_p", "sampling_source", "sampling_status", "sampling_transport"
            ):
                cell.fill = id_fill
            else:  # Le colonne in mezzo sono metriche
                cell.fill = metrics_fill
        else:
            cell.fill = default_fill

    sheet.row_dimensions[1].height = 32
    for column_index, column in enumerate(frame.columns, start=1):
        values = [str(column)] + [str(value) for value in frame[column].dropna().head(100)]
        width = min(max(max(map(len, values), default=12) + 2, 12), 52)
        if column in ("ID", "run_id", "prompt_sha256", "correction_sha256"):
            width = 44 if column in ("ID", "run_id") else 38
        elif column in ("reason", "test_description"):
            width = 52
        sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = width
        is_skill_mode_col = (column == "skill_mode")
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_index, column_index)
            
            # Gestione sfondo: alternato di default, oppure basato su skill_mode
            if is_skill_mode_col:
                val = str(cell.value) if cell.value else ""
                bg_color = SKILL_MODE_COLORS.get(val)
                if bg_color:
                    cell.fill = PatternFill("solid", fgColor=bg_color)
                else:
                    cell.fill = body_fill if row_index % 2 == 0 else PatternFill(fill_type=None)
            else:
                cell.fill = body_fill if row_index % 2 == 0 else PatternFill(fill_type=None)
                
            cell.alignment = Alignment(vertical="top", wrap_text=column in ("reason", "test_description"))
            if column.endswith("rate") or column == "pass_rate":
                cell.number_format = "0.00%"
            elif column == "cost":
                cell.number_format = "$#,##0.0000"
            elif column in ("llm_temperature", "llm_top_p"):
                cell.number_format = "0.0#"
            elif "seconds" in column:
                cell.number_format = "0.0"
            elif "tokens" in column or column in (
                "repetition", "replicate", "run_number",
                "checks_passed", "checks_total", "total_runs",
                "pipeline_completed_runs", "checker_executed_runs", "evaluated_runs",
                "sampling_observed_runs",
                "task_success_runs", "task_failed_runs", "run_count",
                "scenarios", "prompt_types", "skill_modes", "replicates",
                "infrastructure_failures", "completed_runs",
            ):
                cell.number_format = "#,##0"
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 21


def _write_excel(
    path: Path,
    frames: tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame],
    extra_frames: dict[str, pd.DataFrame] | None = None,
) -> None:
    temporary = path.with_name(f"{path.stem}.tmp{path.suffix}")
    all_sheets = list(zip(("Runs", "Checks", "Summary"), frames))
    if extra_frames:
        all_sheets += list(extra_frames.items())
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        writer.book.properties.created = datetime(2000, 1, 1)
        writer.book.properties.modified = datetime(2000, 1, 1)
        for name, frame in all_sheets:
            frame.to_excel(writer, sheet_name=name, index=False, na_rep="")
            _style_sheet(writer.sheets[name], frame, is_runs_sheet=(name == "Runs"))
    canonical = path.with_name(f"{path.stem}.canonical{path.suffix}")
    with zipfile.ZipFile(temporary, "r") as source, zipfile.ZipFile(
        canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = original.create_system
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>", content,
                )
            target.writestr(info, content, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    os.replace(canonical, temporary)
    temporary.replace(path)


def aggregate(runs: Path, results: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Rigenera CSV e workbook dai manifest, metrics e report già salvati."""
    records, check_records = [], []
    for metrics_path in runs.rglob("metrics.json"):
        try:
            parts = metrics_path.relative_to(runs).parts
        except ValueError:
            continue
        if (len(parts) != 6 or parts[4:] != ("evaluation", "metrics.json")
                or not re.fullmatch(r"r\d{3,}", parts[3])):
            continue
        run = metrics_path.parent.parent
        metrics = _read_json(metrics_path)
        manifest_path = run / "manifest.json"
        manifest = _read_json(manifest_path) if manifest_path.is_file() else {}
        record = _run_record(metrics, manifest)
        records.append(record)
        check_records.extend(_check_records(run, record))

    records.sort(key=_sort_key)
    check_records.sort(key=lambda row: (
        _sort_key(row), row.get("category") or "", row.get("test_description") or "",
        bool(row.get("passed")), row.get("reason") or "",
    ))
    run_frame = pd.DataFrame([{key: record.get(key) for key in RUN_COLUMNS} for record in records],
                             columns=RUN_COLUMNS)
    check_frame = pd.DataFrame(check_records, columns=CHECK_COLUMNS)
    summary_records = _summary_records(records)
    summary_frame = pd.DataFrame(summary_records, columns=SUMMARY_COLUMNS)

    results.mkdir(parents=True, exist_ok=True)
    _atomic_csv(results / "runs.csv", run_frame)
    _atomic_csv(results / "checks.csv", check_frame)
    _atomic_csv(results / "summary.csv", summary_frame)
    _write_excel(results / "benchmark.xlsx", (run_frame, check_frame, summary_frame))
    return run_frame, check_frame, summary_frame
