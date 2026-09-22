"""Generazione del report Excel professionale benchmark_report.xlsx con openpyxl.

Crea i fogli:
- Runs: panoramica delle run e punteggi Kathara Lab Checker
- Telemetry: telemetria dettagliata Inspect AI (token, cache, durata, tool calls, errori)
- Checks: dettaglio granulare di ciascun check eseguito dal checker
- Analysis: statistiche aggregate generali e per esperimento
"""
from pathlib import Path
import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from .analysis import summarize

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=False)

SUBHEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E78")

BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

GREEN_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
GREEN_FONT = Font(color="155724", bold=True)
RED_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
RED_FONT = Font(color="721C24", bold=True)


def _autofit_columns(ws, min_width=12, max_width=60):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                val_str = str(val)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def _style_headers_and_filters(ws, freeze_cell="A2"):
    ws.row_dimensions[1].height = 26
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    if freeze_cell:
        ws.freeze_panes = freeze_cell
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions


def _build_runs_sheet(ws, summary_frame: pd.DataFrame):
    runs_columns = [
        ("run_id", "Run ID"),
        ("scenario_id", "Scenario"),
        ("agent", "Agent"),
        ("model", "Model"),
        ("repetition", "Repetition"),
        ("task_success", "Task Success"),
        ("checks_passed", "Checks Passed"),
        ("checks_failed", "Checks Failed"),
        ("checks_total", "Checks Total"),
        ("check_pass_rate", "Check Pass Rate"),
        ("dns_authority_pass_rate", "DNS Auth Rate"),
        ("local_ns_pass_rate", "Local NS Rate"),
        ("dns_record_pass_rate", "DNS Record Rate"),
        ("http_pass_rate", "HTTP Rate"),
        ("reachability_pass_rate", "Reachability Rate"),
        ("files_changed", "Files Changed"),
        ("files_created", "Files Created"),
        ("files_modified", "Files Modified"),
        ("lines_added", "Lines Added"),
        ("lines_deleted", "Lines Deleted"),
        ("pipeline_state", "Pipeline State"),
    ]

    headers = [label for _, label in runs_columns]
    ws.append(headers)

    rate_col_indices = []
    task_success_col = None

    for col_idx, (key, _) in enumerate(runs_columns, start=1):
        if "rate" in key:
            rate_col_indices.append(col_idx)
        if key == "task_success":
            task_success_col = col_idx

    for _, row in summary_frame.iterrows():
        row_vals = []
        for key, _ in runs_columns:
            val = row.get(key)
            if pd.isna(val):
                row_vals.append(None)
            elif key == "task_success":
                # Convert to explicit boolean or string if present
                if isinstance(val, bool):
                    row_vals.append("PASS" if val else "FAIL")
                elif str(val).lower() == "true":
                    row_vals.append("PASS")
                elif str(val).lower() == "false":
                    row_vals.append("FAIL")
                else:
                    row_vals.append(None)
            elif key in ("checks_passed", "checks_failed", "checks_total", "files_changed",
                         "files_created", "files_modified", "lines_added", "lines_deleted", "repetition"):
                try:
                    row_vals.append(int(val))
                except (ValueError, TypeError):
                    row_vals.append(val)
            elif "rate" in key:
                try:
                    row_vals.append(float(val))
                except (ValueError, TypeError):
                    row_vals.append(val)
            else:
                row_vals.append(val)
        ws.append(row_vals)

    # Styling cells
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER_THIN
            if col_idx in rate_col_indices and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
            elif col_idx == task_success_col:
                cell.alignment = Alignment(horizontal="center")

    _style_headers_and_filters(ws)

    # Conditional formatting on task_success
    if task_success_col and ws.max_row > 1:
        col_letter = get_column_letter(task_success_col)
        range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN_FILL, font=GREEN_FONT)
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED_FILL, font=RED_FONT)
        )

    _autofit_columns(ws)


def _build_telemetry_sheet(ws, summary_frame: pd.DataFrame):
    telemetry_columns = [
        ("run_id", "Run ID"),
        ("scenario_id", "Scenario"),
        ("agent", "Agent"),
        ("model", "Model"),
        ("inspect_status", "Inspect Status"),
        ("input_tokens", "Input Tokens"),
        ("output_tokens", "Output Tokens"),
        ("reasoning_tokens", "Reasoning Tokens"),
        ("input_tokens_cache_read", "Cache Read Tokens"),
        ("input_tokens_cache_write", "Cache Write Tokens"),
        ("total_tokens", "Total Tokens"),
        ("total_time", "Duration (s)"),
        ("working_time", "Working Time (s)"),
        ("tool_calls", "Tool Calls"),
        ("tool_errors", "Tool Errors"),
        ("turn_count", "Turns"),
        ("cost", "Cost ($)"),
        ("inspect_log", "Inspect Log File"),
        ("metrics_error", "Metrics Error"),
    ]

    headers = [label for _, label in telemetry_columns]
    ws.append(headers)

    integer_cols = []
    float_cols = []
    status_col = None

    for col_idx, (key, _) in enumerate(telemetry_columns, start=1):
        if key in ("input_tokens", "output_tokens", "reasoning_tokens", "input_tokens_cache_read",
                   "input_tokens_cache_write", "total_tokens", "tool_calls", "tool_errors", "turn_count"):
            integer_cols.append(col_idx)
        elif key in ("total_time", "working_time", "cost"):
            float_cols.append(col_idx)
        elif key == "inspect_status":
            status_col = col_idx

    for _, row in summary_frame.iterrows():
        row_vals = []
        for key, _ in telemetry_columns:
            val = row.get(key)
            if pd.isna(val):
                row_vals.append(None)
            elif key in ("input_tokens", "output_tokens", "reasoning_tokens", "input_tokens_cache_read",
                         "input_tokens_cache_write", "total_tokens", "tool_calls", "tool_errors", "turn_count"):
                try:
                    row_vals.append(int(val))
                except (ValueError, TypeError):
                    row_vals.append(val)
            elif key in ("total_time", "working_time", "cost"):
                try:
                    row_vals.append(float(val))
                except (ValueError, TypeError):
                    row_vals.append(val)
            elif key == "inspect_log" and val:
                try:
                    import os
                    from pathlib import Path
                    ROOT = Path(__file__).resolve().parents[1]
                    p = Path(val).resolve()
                    val = str(p.relative_to(ROOT))
                except ValueError:
                    pass
                row_vals.append(val)
            else:
                row_vals.append(val)
        ws.append(row_vals)

    # Style cells
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER_THIN
            if col_idx in integer_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            elif col_idx in float_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    _style_headers_and_filters(ws)

    # Conditional formatting on inspect_status
    if status_col and ws.max_row > 1:
        col_letter = get_column_letter(status_col)
        range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"success"'], fill=GREEN_FILL, font=GREEN_FONT)
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"error"'], fill=RED_FILL, font=RED_FONT)
        )

    _autofit_columns(ws)


def _build_checks_sheet(ws, detail_frame: pd.DataFrame):
    checks_columns = [
        ("run_id", "Run ID"),
        ("scenario_id", "Scenario"),
        ("repetition", "Repetition"),
        ("agent", "Agent"),
        ("Test Description", "Test Description"),
        ("Passed", "Passed"),
        ("Reason", "Reason / Output"),
    ]

    headers = [label for _, label in checks_columns]
    ws.append(headers)

    passed_col = None
    for col_idx, (key, _) in enumerate(checks_columns, start=1):
        if key == "Passed":
            passed_col = col_idx

    for _, row in detail_frame.iterrows():
        row_vals = []
        for key, _ in checks_columns:
            val = row.get(key)
            if pd.isna(val):
                row_vals.append(None)
            elif key == "Passed":
                if isinstance(val, bool):
                    row_vals.append("PASS" if val else "FAIL")
                elif str(val).lower() in ("true", "pass"):
                    row_vals.append("PASS")
                elif str(val).lower() in ("false", "fail"):
                    row_vals.append("FAIL")
                else:
                    row_vals.append(str(val))
            elif key == "repetition":
                try:
                    row_vals.append(int(val))
                except (ValueError, TypeError):
                    row_vals.append(val)
            else:
                row_vals.append(val)
        ws.append(row_vals)

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 19
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER_THIN
            if col_idx == passed_col:
                cell.alignment = Alignment(horizontal="center")

    _style_headers_and_filters(ws)

    # Conditional formatting on Passed
    if passed_col and ws.max_row > 1:
        col_letter = get_column_letter(passed_col)
        range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN_FILL, font=GREEN_FONT)
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED_FILL, font=RED_FONT)
        )

    _autofit_columns(ws)


def _build_analysis_sheet(ws, summary_frame: pd.DataFrame):
    ws.views.sheetView[0].showGridLines = True
    stats = summarize(summary_frame) if not summary_frame.empty else {}

    # Section 1: Executive KPI Table
    ws.append(["BENCHMARK OVERALL SUMMARY", ""])
    ws.merge_cells("A1:B1")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = HEADER_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    kpi_rows = [
        ("Total Runs", stats.get("runs", 0), "integer"),
        ("Evaluated Runs (Checker Executed)", stats.get("evaluated_runs", 0), "integer"),
        ("Unscored / Incomplete Runs", stats.get("unscored_runs", 0), "integer"),
        ("Task Success Rate", stats.get("task_success_rate"), "percentage"),
        ("Mean Check Pass Rate", stats.get("check_pass_rate_mean"), "percentage"),
        ("Median Check Pass Rate", stats.get("check_pass_rate_median"), "percentage"),
        ("Mean Total Tokens", stats.get("total_tokens_mean"), "integer"),
        ("Median Total Tokens", stats.get("total_tokens_median"), "integer"),
        ("Mean Duration (seconds)", stats.get("total_time_mean"), "float"),
        ("Median Duration (seconds)", stats.get("total_time_median"), "float"),
        ("Mean Working Time (seconds)", stats.get("working_time_mean"), "float"),
        ("Mean Tool Calls", stats.get("tool_calls_mean"), "float"),
        ("Tool Error Rate", stats.get("tool_error_rate"), "percentage_4"),
        ("Mean Files Changed", stats.get("files_changed_mean"), "float"),
        ("Mean Lines Added", stats.get("lines_added_mean"), "float"),
    ]

    for label, val, kind in kpi_rows:
        ws.append([label, val])
        row_idx = ws.max_row
        ws.row_dimensions[row_idx].height = 20
        cell_lbl = ws.cell(row=row_idx, column=1)
        cell_val = ws.cell(row=row_idx, column=2)
        cell_lbl.border = BORDER_THIN
        cell_val.border = BORDER_THIN
        cell_lbl.font = Font(name="Calibri", size=11, bold=True)
        cell_lbl.fill = SUBHEADER_FILL
        if val is not None and isinstance(val, (int, float)):
            if kind == "integer":
                cell_val.number_format = "#,##0"
            elif kind == "percentage":
                cell_val.number_format = "0.0%"
            elif kind == "percentage_4":
                cell_val.number_format = "0.00%"
            elif kind == "float":
                cell_val.number_format = "#,##0.00"

    ws.append([])  # Blank row

    # Section 2: Experiment Breakdown
    breakdown_header_row = ws.max_row + 1
    breakdown_headers = [
        "Scenario", "Agent", "Model", "Runs", "Evaluated", "Task Success Rate",
        "Check Pass Rate (Mean)", "Total Tokens (Mean)", "Duration Mean (s)",
        "Tool Calls (Mean)", "Tool Error Rate"
    ]
    ws.append(breakdown_headers)
    ws.row_dimensions[breakdown_header_row].height = 26
    for col_idx in range(1, len(breakdown_headers) + 1):
        cell = ws.cell(row=breakdown_header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    if not summary_frame.empty:
        grouped = summary_frame.groupby(["scenario_id", "agent", "model"], dropna=False)
        for (scenario, agent, model), group in grouped:
            g_stats = summarize(group)
            ws.append([
                scenario,
                agent,
                model,
                g_stats.get("runs", 0),
                g_stats.get("evaluated_runs", 0),
                g_stats.get("task_success_rate"),
                g_stats.get("check_pass_rate_mean"),
                g_stats.get("total_tokens_mean"),
                g_stats.get("total_time_mean"),
                g_stats.get("tool_calls_mean"),
                g_stats.get("tool_error_rate"),
            ])
            r_idx = ws.max_row
            ws.row_dimensions[r_idx].height = 20
            for c_idx in range(1, len(breakdown_headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = BORDER_THIN
                if c_idx in (6, 7) and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0%"
                elif c_idx == 11 and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00%"
                elif c_idx == 8 and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
                elif c_idx in (9, 10) and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

    _autofit_columns(ws)


def generate_excel_report(results: Path) -> Path:
    """Genera results/benchmark_report.xlsx aggregando Runs, Telemetry, Checks e Analysis dai file CSV."""
    summary_csv = results / "benchmark_results.csv"
    detail_csv = results / "benchmark_detailed.csv"
    
    if not summary_csv.exists() or not detail_csv.exists():
        return None
        
    summary_frame = pd.read_csv(summary_csv)
    detail_frame = pd.read_csv(detail_csv)
    
    results.mkdir(parents=True, exist_ok=True)
    report_path = results / "benchmark_report.xlsx"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # 1. Runs Sheet
    ws_runs = wb.create_sheet(title="Runs")
    _build_runs_sheet(ws_runs, summary_frame)

    # 2. Telemetry Sheet
    ws_telemetry = wb.create_sheet(title="Telemetry")
    _build_telemetry_sheet(ws_telemetry, summary_frame)

    # 3. Checks Sheet
    ws_checks = wb.create_sheet(title="Checks")
    _build_checks_sheet(ws_checks, detail_frame)

    # 4. Analysis Sheet
    ws_analysis = wb.create_sheet(title="Analysis")
    _build_analysis_sheet(ws_analysis, summary_frame)

    # Salva in modo atomico
    temp_path = results / "benchmark_report.xlsx.tmp"
    wb.save(temp_path)
    temp_path.replace(report_path)
    return report_path
